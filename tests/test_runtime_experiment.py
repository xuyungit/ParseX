"""The agent's working directory (plan P2-1, §2.2–§2.5; Q85): the px wrapper, the task, post-run verification."""

import json

import pytest

from parserx.config.schema import ParserXConfig
from parserx.runtimes.experiment import doc_config, render_task, secret_names, verify_run
from parserx.runtimes.px import PxRefused, px_argv, python_env
from parserx.tools import call_tool, workspace_init

RAW = {
    "services": {"vlm": {"endpoint": "${OPENAI_BASE_URL_B}", "api_key": "${OPENAI_API_KEY_B}",
                         "model": "${VLM_MODEL_B:gpt-6-luna}"}},
    "builders": {"ocr": {"token": "${PADDLE_OCR_TOKEN}"}},
    "cache": {"mode": "read_write", "dir": ".parserx_cache"},
}


def test_a_run_config_keeps_secrets_as_placeholders_and_the_cache_inside(tmp_path):
    config = doc_config(RAW, tmp_path / "run")
    assert config["services"]["vlm"]["api_key"] == "${OPENAI_API_KEY_B}"
    assert config["cache"] == {"mode": "read_write", "dir": str(tmp_path / "run" / ".parserx_cache")}
    assert secret_names(RAW) == ["OPENAI_API_KEY_B", "OPENAI_BASE_URL_B", "PADDLE_OCR_TOKEN", "VLM_MODEL_B"]


def test_px_fixes_the_config_and_allows_only_the_agents_tools(tmp_path):
    cfg = tmp_path / "parserx.yaml"
    assert px_argv(["tool", "read_draft", "--ws", "ws", "--view", "issues", "--json"], cfg) == [
        "tool", "read_draft", "--ws", "ws", "--view", "issues", "--json", "--config", str(cfg)]
    assert px_argv(["tool", "schema", "edit_draft"], cfg) == ["tool", "schema", "edit_draft"]
    for refused in (["parse", "input.pdf"], ["eval", "gt"], ["tool", "read_draft", "--config", "x.yaml"],
                    ["tool", "read_draft", "-c", "x.yaml"], ["tool", "frobnicate"], ["tool", "run_pipeline"],
                    ["workspace", "init", "input.pdf", "--ws", "ws"]):
        with pytest.raises(PxRefused):
            px_argv(refused, cfg)


def test_px_python_runs_without_service_keys():
    env = python_env({"PATH": "/usr/bin", "OPENAI_API_KEY_B": "sk", "PADDLE_OCR_TOKEN": "t", "MY_SECRET": "s",
                      "OPENAI_BASE_URL_B": "https://x", "HOME": "/h"}, secret_names={"OPENAI_BASE_URL_B"})
    assert env == {"PATH": "/usr/bin", "HOME": "/h"}


def _docx(path):
    import docx

    document = docx.Document()
    document.add_heading("SENTINEL 标题", level=1)
    document.add_paragraph("SENTINEL 正文 100 万元")
    document.save(path)
    return path


def test_verify_run_checks_integrity_and_counts_tool_calls(tmp_path):
    doc_dir = tmp_path / "doc"
    doc_dir.mkdir()
    source = _docx(doc_dir / "input.docx")
    config = ParserXConfig()
    workspace_init(source, doc_dir / "ws", config=config)
    call_tool("read_draft", doc_dir / "ws", {}, config=config)
    call_tool("read_draft", doc_dir / "ws", {"view": "blocks", "blocks": ["b-nope"]}, config=config)  # a failed call
    env, _ = call_tool("submit_draft", doc_dir / "ws", {}, config=config)
    assert env.ok and env.result.accepted

    report = verify_run(doc_dir)
    assert report.integrity.ok
    assert report.tools.calls == {"read_draft": 2, "submit_draft": 1, "workspace_init": 1}
    assert report.tools.failures == {"not_found": 1}
    assert report.check.exportable and report.check.document_status == "complete"


def test_verify_run_reports_direct_changes(tmp_path):
    doc_dir = tmp_path / "doc"
    doc_dir.mkdir()
    source = _docx(doc_dir / "input.docx")
    workspace_init(source, doc_dir / "ws", config=ParserXConfig())
    raw = json.loads((doc_dir / "ws" / "state.json").read_text())
    raw["blocks"][1]["text"] = "改写"
    (doc_dir / "ws" / "state.json").write_text(json.dumps(raw, ensure_ascii=False))
    assert not verify_run(doc_dir).integrity.ok


def test_task_template_renders_every_value_and_the_chosen_blocks():
    template = "# {{input_name}}\n{{#vision_agent}}look yourself\n{{/vision_agent}}{{#vision_tool}}ask the tool\n" \
               "{{/vision_tool}}budget {{minutes}}\n"
    values = {"input_name": "input.pdf", "minutes": 30}
    assert render_task(template, round_name=None, values=values, options={"vision_tool"}) == \
        "# input.pdf\nask the tool\nbudget 30\n"
    assert render_task(template, round_name=None, values=values, options={"vision_agent"}) == \
        "# input.pdf\nlook yourself\nbudget 30\n"
    with pytest.raises(KeyError):
        render_task(template, round_name=None, values={"input_name": "input.pdf"})


def test_the_shipped_task_is_the_task_the_cli_adapter_and_the_generated_reference():
    from parserx.runtimes.experiment import compose_task

    tool = compose_task("cli", input_name="input.pdf", minutes=30, vision="tool")
    agent = compose_task("cli", input_name="input.pdf", minutes=30, vision="agent")
    for text in (tool, agent):
        assert "{{" not in text and "input.pdf" in text and "30 分钟" in text and "./px tool" in text
        for name in ("read_draft", "view_source", "edit_draft", "submit_draft", "doc_text", "evidence"):
            assert name in text
        for op in ("replace_text", "set_cells", "insert_text", "adopt", "set_role", "move", "join", "unjoin", "split",
                   "exclude", "include", "mark_pending", "dismiss", "note"):  # every operation, from the schema
            assert f"**`{op}`**" in text, op
        for old in ("./px tool process", "overview", "ask_image", "apply_structure", "review_table", "describe_figure",
                    "workspace init", "`correct`", "`close`", "`skim`", "set_level", "merge_tables", '"link"',
                    "--from", "--out"):
            assert old not in text, old
    # the two ways of seeing images differ only where the source is looked at
    assert "打开这个文件亲自看" in agent and "不要用 `as: image`" not in agent
    assert "不要用 `as: image`" in tool and "打开这个文件亲自看" not in tool
    # without the command line, the task says nothing of it
    assert "./px" not in compose_task("call", input_name="input.pdf", minutes=30)
