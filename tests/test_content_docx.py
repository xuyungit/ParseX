"""DOCX reader (guide §6.9, plan R1/R3, Q26) on synthetic documents: python-docx plus hand-written XML."""

import io

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from PIL import Image

from parserx.content.docx import extract_docx
from parserx.ir.anchor import AssetAnchor, DocxAnchor
from parserx.ir.enums import BlockKind, BlockStatus

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _add(body, fragment: str) -> None:
    """Insert a paragraph where python-docx would: before the body-level sectPr."""
    body[-1].addprevious(parse_xml(fragment))


def _png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), (10, 120, 200)).save(buf, "PNG")
    return buf.getvalue()


@pytest.fixture
def docx_path(tmp_path):
    doc = Document()
    doc.add_paragraph("Annual Report", style="Title")
    doc.add_paragraph("Overview", style="Heading 1")
    body = doc.element.body
    # tracked changes: final view keeps insertions and drops deletions
    _add(body, (
        f'<w:p {W}><w:r><w:t xml:space="preserve">Revenue was </w:t></w:r>'
        '<w:del w:id="1" w:author="a"><w:r><w:delText>90</w:delText></w:r></w:del>'
        '<w:ins w:id="2" w:author="a"><w:r><w:t>100</w:t></w:r></w:ins>'
        '<w:r><w:t xml:space="preserve"> million.</w:t></w:r></w:p>'))
    _add(body, (
        f'<w:p {W}><w:moveFrom w:id="3" w:author="a"><w:r><w:t>Moved away.</w:t></w:r></w:moveFrom>'
        '<w:r><w:t>Stays.</w:t></w:r></w:p>'))
    # a complex field: only its result text is content
    _add(body, (
        f'<w:p {W}><w:r><w:t xml:space="preserve">See page </w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> PAGEREF _Toc1 </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>7</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        '<w:fldSimple w:instr=" DATE "><w:r><w:t xml:space="preserve"> (2026)</w:t></w:r></w:fldSimple></w:p>'))
    # a direct outline level on a plain paragraph
    _add(body, (f'<w:p {W}><w:pPr><w:outlineLvl w:val="1"/></w:pPr><w:r><w:t>Direct level</w:t></w:r></w:p>'))
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged header"
    table.cell(0, 2).text = "C"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Tall"
    for r, c, v in [(1, 1, "10"), (1, 2, "11"), (2, 1, "20"), (2, 2, "21")]:
        table.cell(r, c).text = v
    doc.add_page_break()
    doc.add_paragraph("After the page break")
    image_stream = io.BytesIO(_png())
    doc.add_picture(image_stream)
    doc.add_section(WD_SECTION.CONTINUOUS)
    doc.add_paragraph("New section")
    # a textbox (Q44): its paragraphs become text after the paragraph that anchors it
    _add(body, (
        f'<w:p {W} xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><mc:AlternateContent><mc:Choice Requires="wps">'
        '<w:drawing><wps:txbx><w:txbxContent><w:p><w:r><w:t>Callout text</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Second line</w:t></w:r></w:p></w:txbxContent>'
        '</wps:txbx></w:drawing></mc:Choice><mc:Fallback><w:pict><v:textbox><w:txbxContent><w:p><w:r>'
        '<w:t>Callout text</w:t></w:r></w:p></w:txbxContent></v:textbox></w:pict></mc:Fallback>'
        '</mc:AlternateContent></w:r></w:p>'))
    path = tmp_path / "doc.docx"
    doc.save(path)
    return path


def _text_blocks(ext):
    return [b for b in ext.blocks if b.kind in (BlockKind.TEXT, BlockKind.TITLE) and b.status == BlockStatus.OK]


def test_final_view_of_tracked_changes(docx_path):
    ext = extract_docx(docx_path)
    texts = [b.text for b in _text_blocks(ext)]
    assert "Revenue was 100 million." in texts and "Stays." in texts
    deleted = [b for b in ext.blocks if b.status == BlockStatus.EXCLUDED]
    assert sorted(b.text for b in deleted) == ["90", "Moved away."]
    assert all(b.decisions[-1].stage == "exclude" and b.decisions[-1].choice == "revision_deleted" for b in deleted)
    ledger = {e.block: e for e in ext.ledger}
    assert all(ledger[b.id].unit == "docx_deleted" and ledger[b.id].disposition == "excluded" for b in deleted)
    assert any("tracked changes" in w for w in ext.warnings)


def test_fields_keep_only_their_results(docx_path):
    texts = [b.text for b in _text_blocks(extract_docx(docx_path))]
    assert "See page 7 (2026)" in texts
    assert not any("PAGEREF" in t or "DATE" in t for t in texts)


def test_style_and_outline_evidence(docx_path):
    blocks = {b.text: b for b in _text_blocks(extract_docx(docx_path))}
    title = blocks["Annual Report"].observations[0].style
    heading = blocks["Overview"].observations[0].style
    direct = blocks["Direct level"].observations[0].style
    assert title.style_name == "Title"
    assert heading.style_name == "heading 1" and heading.outline_level == 0  # built-in names are stored lowercase
    assert direct.outline_level == 1 and direct.style_name == "Normal"
    assert blocks["Overview"].kind == BlockKind.TEXT  # roles are the structure step's job


def test_table_spans(docx_path):
    table = next(b for b in extract_docx(docx_path).blocks if b.kind == BlockKind.TABLE)
    grid = table.cells
    assert (grid.n_rows, grid.n_cols) == (3, 3)
    head = grid.slot(0, 1)
    assert head.content == "Merged header" and head.colspan == 2
    tall = grid.slot(2, 0)
    assert tall.content == "Tall" and tall.rowspan == 2
    assert grid.slot(2, 2).content == "21"


def test_segments_follow_page_and_section_breaks(docx_path):
    ext = extract_docx(docx_path)
    seg = {b.text: b.anchors[0].segment for b in _text_blocks(ext)}
    assert seg["Overview"] == 1 and seg["After the page break"] == 2 and seg["New section"] == 3
    assert [(p.n, p.unit, p.starts_with) for p in ext.pages] == [
        (1, "docx_segment", None), (2, "docx_segment", "page_break"), (3, "docx_segment", "section_break")]


def test_images_become_assets(docx_path):
    ext = extract_docx(docx_path)
    figure = next(b for b in ext.blocks if b.kind == BlockKind.FIGURE)
    asset_anchor = next(a for a in figure.anchors if isinstance(a, AssetAnchor))
    asset = next(a for a in ext.assets if a.id == asset_anchor.asset)
    assert (asset.width, asset.height) == (40, 30) and asset.source.part == "word/document.xml"
    assert figure.anchors[0].segment == 2


def test_textbox_paragraphs_become_text_after_their_anchor(docx_path):
    ext = extract_docx(docx_path)
    texts = [b.text for b in _text_blocks(ext)]
    assert texts[texts.index("New section") + 1:] == ["Callout text", "Second line"]  # the VML fallback is not read twice
    assert not any(b.status == BlockStatus.FAILED for b in ext.blocks)
    boxes = [b for b in ext.blocks if b.text in ("Callout text", "Second line")]
    entries = [e for e in ext.ledger if e.block in {b.id for b in boxes}]
    assert [(e.unit, e.disposition, e.chars) for e in entries] == [("docx_paragraph", "output", 11),
                                                                    ("docx_paragraph", "output", 10)]
    assert boxes[0].anchors[0].node_path.endswith("w:txbxContent[1]/w:p[1]")
    assert not any("textbox" in w for w in ext.warnings)


def test_ledger_and_anchors(docx_path):
    ext = extract_docx(docx_path)
    assert all(e.disposition is not None and e.block for e in ext.ledger)
    assert [e.item for e in ext.ledger] == [f"i-d{i:05d}" for i in range(1, len(ext.ledger) + 1)]
    anchor = _text_blocks(ext)[0].anchors[0]
    assert isinstance(anchor, DocxAnchor) and anchor.part == "word/document.xml"
    assert anchor.node_path == "/w:body/w:p[1]"
    assert [b.id for b in ext.blocks] == [f"b-d{i:05d}" for i in range(1, len(ext.blocks) + 1)]


def test_deleted_paragraph_mark_merges_paragraphs(tmp_path):
    doc = Document()
    body = doc.element.body
    _add(body, (
        f'<w:p {W}><w:pPr><w:rPr><w:del w:id="9" w:author="a"/></w:rPr></w:pPr>'
        '<w:r><w:t xml:space="preserve">First half, </w:t></w:r></w:p>'))
    _add(body, (f'<w:p {W}><w:r><w:t>second half.</w:t></w:r></w:p>'))
    path = tmp_path / "merge.docx"
    doc.save(path)
    ext = extract_docx(path)
    assert [b.text for b in _text_blocks(ext)] == ["First half, second half."]
    merged = [e for e in ext.ledger if e.disposition == "merged"]
    assert len(merged) == 1 and merged[0].block == _text_blocks(ext)[0].id


def test_list_numbering_is_rendered(tmp_path):
    doc = Document()
    for text in ("alpha", "beta"):
        doc.add_paragraph(text, style="List Number")
    path = tmp_path / "list.docx"
    doc.save(path)
    blocks = _text_blocks(extract_docx(path))
    assert [b.text for b in blocks] == ["1. alpha", "2. beta"]
    numbering = blocks[1].observations[0].style.numbering
    assert numbering.level == 0 and numbering.text == "2."


def test_extraction_is_deterministic(docx_path):
    a, b = extract_docx(docx_path), extract_docx(docx_path)
    assert [x.model_dump() for x in a.blocks] == [x.model_dump() for x in b.blocks] and a.ledger == b.ledger


def test_charts_and_smartart_are_accounted_for_with_their_text(tmp_path):
    # P4-5 / Q9: a drawing holding a chart or SmartArt has no image; its content lives in a part of its own.
    # Not read into the output yet, but never dropped silently: a failed block with the part's text, a warning.
    import zipfile

    doc = Document()
    doc.add_paragraph("Figure 1 shows the output.")
    body = doc.element.body
    graphic = ('<w:p {W} xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
               'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
               'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:r><w:drawing>'
               '<wp:inline><wp:extent cx="4000000" cy="3000000"/><a:graphic><a:graphicData uri="{uri}">{ref}'
               '</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>')
    _add(body, graphic.format(W=W, uri="http://schemas.openxmlformats.org/drawingml/2006/chart",
                              ref='<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
                                  'r:id="rIdChart1"/>'))
    _add(body, graphic.format(W=W, uri="http://schemas.openxmlformats.org/drawingml/2006/diagram",
                              ref='<dgm:relIds xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
                                  'r:dm="rIdDm1" r:lo="rIdLo1"/>'))
    plain = tmp_path / "plain.docx"
    doc.save(plain)
    chart = ('<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
             'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:chart><c:title><c:tx><c:rich><a:p>'
             '<a:r><a:t>月产量</a:t></a:r></a:p></c:rich></c:tx></c:title><c:plotArea><c:barChart><c:ser>'
             '<c:cat><c:strRef><c:strCache><c:pt idx="0"><c:v>一月</c:v></c:pt></c:strCache></c:strRef></c:cat>'
             '<c:val><c:numRef><c:numCache><c:pt idx="0"><c:v>120</c:v></c:pt></c:numCache></c:numRef></c:val>'
             '</c:ser></c:barChart></c:plotArea></c:chart></c:chartSpace>')
    data = ('<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dgm:ptLst><dgm:pt><dgm:t><a:p><a:r>'
            '<a:t>需求分析</a:t></a:r></a:p></dgm:t></dgm:pt></dgm:ptLst></dgm:dataModel>')
    path = tmp_path / "drawings.docx"
    with zipfile.ZipFile(plain) as src, zipfile.ZipFile(path, "w") as out:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                content = content.replace(b"</Relationships>", (
                    '<Relationship Id="rIdChart1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    'relationships/chart" Target="charts/chart1.xml"/>'
                    '<Relationship Id="rIdDm1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                    'relationships/diagramData" Target="diagrams/data1.xml"/></Relationships>').encode())
            out.writestr(item, content)
        out.writestr("word/charts/chart1.xml", chart)
        out.writestr("word/diagrams/data1.xml", data)
    ext = extract_docx(path)
    failed = {b.text: b for b in ext.blocks if b.status == BlockStatus.FAILED}
    assert set(failed) == {"月产量\n一月\n120", "需求分析"}
    assert all(b.kind == BlockKind.FIGURE for b in failed.values())
    assert len(ext.missing) == 2
    assert {e.disposition for e in ext.ledger if e.block in {b.id for b in failed.values()}} == {"failed"}


def _with_parts(plain, path, parts: dict[str, str], rels: list[tuple[str, str, str]]):
    import zipfile

    with zipfile.ZipFile(plain) as src, zipfile.ZipFile(path, "w") as out:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                extra = "".join(f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/'
                                f'2006/relationships/{kind}" Target="{target}"/>' for rid, kind, target in rels)
                content = content.replace(b"</Relationships>", (extra + "</Relationships>").encode())
            out.writestr(item, content)
        for name, xml in parts.items():
            out.writestr(name, xml)


def test_footnotes_become_markdown_footnotes_and_comments_are_excluded(tmp_path):
    # Q9: the reference stays where Word shows its number; the note follows its paragraph; a comment is a
    # reviewer's note, not content
    from parserx.render import render_markdown

    doc = Document()
    body = doc.element.body
    _add(body, (f'<w:p {W}><w:r><w:t>Load tests</w:t></w:r><w:r><w:footnoteReference w:id="1"/></w:r>'
                '<w:r><w:t xml:space="preserve"> were repeated</w:t></w:r><w:r><w:endnoteReference w:id="2"/></w:r>'
                '<w:commentRangeStart w:id="0"/><w:r><w:t>.</w:t></w:r><w:commentRangeEnd w:id="0"/>'
                '<w:r><w:commentReference w:id="0"/></w:r></w:p>'))
    _add(body, f'<w:p {W}><w:r><w:t>Next paragraph.</w:t></w:r></w:p>')
    plain = tmp_path / "plain.docx"
    doc.save(plain)
    notes = (f'<w:{{kind}}s {W}><w:{{kind}} w:id="0"><w:p><w:r><w:t>separator</w:t></w:r></w:p></w:{{kind}}>'
             f'<w:{{kind}} w:id="{{id}}"><w:p><w:r><w:{{kind}}Ref/></w:r><w:r><w:t>{{text}}</w:t></w:r></w:p>'
             f'</w:{{kind}}></w:{{kind}}s>')
    comments = f'<w:comments {W}><w:comment w:id="0" w:author="r"><w:p><w:r><w:t>Check this</w:t></w:r></w:p></w:comment></w:comments>'
    path = tmp_path / "notes.docx"
    _with_parts(plain, path, {
        "word/footnotes.xml": notes.format(kind="footnote", id=1, text="At 20 °C."),
        "word/endnotes.xml": notes.format(kind="endnote", id=2, text="See annex B."),
        "word/comments.xml": comments,
    }, [("rIdF", "footnotes", "footnotes.xml"), ("rIdE", "endnotes", "endnotes.xml"),
        ("rIdC", "comments", "comments.xml")])
    ext = extract_docx(path)
    assert not ext.missing
    state = ext.to_state(doc_id="notes", source="notes.docx", source_sha256="0" * 64)
    assert render_markdown(state) == ("Load tests[^1] were repeated[^e2].\n\n[^1]: At 20 °C.\n\n[^e2]: See annex B.\n\n"
                                      "Next paragraph.\n")
    comment = next(b for b in ext.blocks if b.text == "Check this")
    assert comment.status == BlockStatus.EXCLUDED and comment.decisions[-1].choice == "comment"
    assert {e.disposition for e in ext.ledger if e.block == comment.id} == {"excluded"}


def test_the_face_a_paragraph_is_set_in_is_style_evidence(tmp_path):
    # Chinese official documents set headings in 黑体 at the body's size: the face is the only typographic difference
    doc = Document()
    defaults = doc.styles.element.find(".//{*}docDefaults")
    defaults.remove(defaults.find("{*}rPrDefault"))
    defaults.append(parse_xml(  # defaults: 仿宋 for CJK text, Times New Roman otherwise
        f'<w:rPrDefault {W}><w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" '
        'w:eastAsia="仿宋_GB2312"/><w:sz w:val="32"/></w:rPr></w:rPrDefault>'))
    heading = doc.add_paragraph()
    run = heading.add_run("一、企业基本情况")
    run._r.get_or_add_rPr().append(parse_xml(f'<w:rFonts {W} w:ascii="Times New Roman" w:eastAsia="黑体"/>'))
    doc.add_paragraph("企业成立于二〇一〇年，主要从事桥梁支座的研发与生产。")
    doc.add_paragraph("Bearing test report")
    path = tmp_path / "faces.docx"
    doc.save(path)
    styles = {b.text: b.observations[0].style for b in _text_blocks(extract_docx(path))}
    assert styles["一、企业基本情况"].font == "黑体"
    assert styles["企业成立于二〇一〇年，主要从事桥梁支座的研发与生产。"].font == "仿宋_GB2312"  # the default
    assert styles["Bearing test report"].font == "Times New Roman"  # Latin text: the ASCII face
