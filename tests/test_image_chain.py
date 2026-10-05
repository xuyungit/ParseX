"""Figures turned, routed and described each as soon as the step before is done (tools/image_chain.py, speed plan
P3): answers that come back in any order are recorded on their own figures, in reading order."""

import io
import json
import threading
import time
from pathlib import Path

import docx
from PIL import Image

from parserx.config.schema import ParserXConfig
from parserx.ir.enums import BlockKind
from parserx.tools.context import ToolContext
from parserx.tools import call_tool
from parserx.tools.init import workspace_init
from parserx.workspace.store import Workspace


class _Reader:
    name, version = "reading", "fake-reader-1"

    def read(self, png):
        return []


class _VLM:
    """Describes an image by its file name; the first request answers last."""

    def __init__(self):
        self.first = threading.Event()

    def describe_image(self, image_path, prompt, **kwargs):
        if not self.first.is_set():
            self.first.set()
            time.sleep(0.3)
        return json.dumps({"type": "photo", "caption": f"照片 {Path(image_path).stem}"})


def test_answers_in_any_order_land_on_their_own_figures(tmp_path):
    d = docx.Document()
    for colour in ("red", "green", "blue"):
        buf = io.BytesIO()
        Image.new("RGB", (40, 30), colour).save(buf, "PNG")
        buf.seek(0)
        d.add_paragraph().add_run().add_picture(buf)
        d.add_paragraph(f"{colour} 下方的说明文字")
    source = tmp_path / "three.docx"
    d.save(source)
    config = ParserXConfig()
    config.cache.mode = "off"
    config.services.vlm.model, config.services.vlm.endpoint, config.services.vlm.api_key = "m", "https://x", "k"
    config.runtime.layout_shadow = False  # no layout model in tests
    vlm = _VLM()

    class Context(ToolContext):
        def _new_reader(self):
            return _Reader()

        def _new_vlm(self, cfg):
            return vlm

    ws = tmp_path / "ws"
    assert workspace_init(source, ws, config=config)[1] == 0
    envelope, code = call_tool("run_pipeline", ws, {}, config=config, context_factory=Context)
    assert code == 0 and "images" in [s.step for s in envelope.result.steps]
    state = Workspace.open(ws).load()
    assets = {a.id: a for a in state.assets}
    figures = [b for b in sorted(state.blocks, key=lambda b: b.order) if b.kind == BlockKind.FIGURE]
    assert len(figures) == 3
    for figure in figures:
        asset = assets[next(a.asset for a in figure.anchors if hasattr(a, "asset"))]
        assert figure.semantic.caption == f"照片 {Path(asset.path).stem}"
