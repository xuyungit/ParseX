"""Inline bold and underline (R3): marks from Word runs and bold PDF spans, written around their spans."""

import pymupdf
from docx import Document

from parserx.content.docx import extract_docx
from parserx.content.pdf_native import extract_pdf
from parserx.ir.observation import Mark
from parserx.render.emphasis import emphasize


def test_marks_are_written_around_their_spans_with_punctuation_outside_bold():
    assert emphasize("项目负责人：邹贻军", [Mark(kind="underline", text="邹贻军")])[0] == "项目负责人：<u>邹贻军</u>"
    assert emphasize("发件人：Apple", [Mark(kind="bold", text="发件人：")])[0] == "**发件人**：Apple"
    assert emphasize("全段粗体", [Mark(kind="bold", text="全段粗体"), Mark(kind="underline", text="全段粗体")])[0] == \
        "**<u>全段粗体</u>**"
    assert emphasize("注 *1 与 重点", [Mark(kind="bold", text="重点")])[0] == "注 \\*1 与 **重点**"
    text, left = emphasize("找不到", [Mark(kind="bold", text="别处")])
    assert text == "找不到" and [m.text for m in left] == ["别处"]


def test_word_runs_give_bold_and_underline_spans(tmp_path):
    doc = Document()
    para = doc.add_paragraph("买方：")
    para.add_run("某工程局").underline = True
    para.add_run("，签订日期 ")
    para.add_run("2025").underline = True
    para.add_run(" 年；")
    para.add_run("注意事项").bold = True
    path = tmp_path / "marks.docx"
    doc.save(path)
    marks = [(m.kind, m.text) for b in extract_docx(path).blocks for o in b.observations for m in o.marks]
    assert sorted(marks) == [("bold", "注意事项"), ("underline", "2025"), ("underline", "某工程局")]


def test_bold_spans_of_native_text_are_marks(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 100), "Order:", fontsize=11, fontname="hebo")
    page.insert_text((115, 100), "MV6LY51NQF placed today", fontsize=11, fontname="helv")
    path = tmp_path / "bold.pdf"
    doc.save(path)
    marks = [(m.kind, m.text) for b in extract_pdf(path).blocks for o in b.observations for m in o.marks]
    assert marks == [("bold", "Order:")]
