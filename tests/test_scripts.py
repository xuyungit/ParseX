"""Sub- and superscripts of the native text layer (content/scripts.py): judged from size and baseline, written as
marks of a block's text and into table cells."""

import pymupdf

from parserx.content import scripts
from parserx.content.pdf_native import extract_pdf
from parserx.ir.enums import BlockKind
from parserx.ir.observation import Mark
from parserx.render.emphasis import emphasize


def _line(*glyphs, main=None):
    """A text-layer line from (char, x0, x1, size, baseline); its main span is *main* (baseline, size), else its
    largest glyph."""
    out = tuple(scripts.Glyph(c, x0, x1, base - 0.8 * size, base + 0.2 * size, size, base) for c, x0, x1, size, base in glyphs)
    if main is None:
        largest = max(out, key=lambda g: g.size)
        main = (largest.baseline, largest.size)
    return scripts.Line(out, *main)


def _row(text, x0, size, base, width=None):
    width = width or size
    return [(c, x0 + k * width, x0 + (k + 1) * width, size, base) for k, c in enumerate(text)]


# ── judging (candidates version 2 of the vision-first branch, eval_reports/2026-09-30_script_candidates.md) ──


def test_a_small_raised_glyph_beside_its_base_is_a_superscript():
    assert scripts.kinds([_line(("x", 0, 5, 10, 100), ("2", 5, 8, 6, 96))]) == [["", "sup"]]
    assert scripts.kinds([_line(("k", 0, 5, 10, 100), ("1", 5, 8, 6, 103))]) == [["", "sub"]]


def test_a_script_kept_as_a_line_of_its_own_is_judged_in_its_row():
    assert scripts.kinds([_line(("m", 0, 8, 10, 100)), _line(("2", 8, 11, 6, 96))]) == [[""], ["sup"]]


def test_glyphs_of_one_size_are_no_scripts():
    assert scripts.kinds([_line(*_row("mc2", 0, 10, 100))]) == [["", "", ""]]


def test_a_smaller_line_under_a_larger_one_is_not_its_subscript():
    found = scripts.kinds([_line(*_row("结构横断面", 0, 18, 100)), _line(*_row("Section", 5, 8.8, 110, 10))])
    assert found == [[""] * 5, [""] * 7]


def test_a_raised_citation_is_one_script_and_the_full_stop_after_it_is_none():
    found = scripts.kinds([_line(("梁", 0, 21, 21, 100), ("[", 21, 27, 12, 95), ("1", 27, 30, 6, 95),
                                 ("]", 30, 36, 12, 95), (".", 36, 39, 6, 100), ("等", 39, 60, 21, 100))])
    assert found == [["", "sup", "sup", "sup", "", ""]]


def test_a_prime_is_a_character_unless_it_stands_in_a_script():
    alone = scripts.kinds([_line(("z", 0, 5, 10, 100), ("′", 5, 8, 5, 96), ("E", 8, 14, 10, 100))])
    inside = scripts.kinds([_line(("q", 0, 5, 10, 100), ("x", 5, 8, 6, 96), ("′", 8, 10, 5, 96))])
    assert alone == [["", "", ""]]
    assert inside == [["", "sup", "sup"]]


def test_an_opening_bracket_carries_no_script():
    # δ = ( … in a display formula: the parenthesis is set larger than the formula's text, not its scripts' base
    found = scripts.kinds([_line(("δ", 0, 6, 10.3, 114.1), ("=", 13, 19, 10.3, 113.6), ("(", 23, 27, 14.6, 116.2),
                                 main=(114.1, 10.3))])
    assert found == [["", "", ""]]


def test_the_text_beside_an_enlarged_operator_is_no_script_of_it():
    # ∑_{k=1} δ: the limit is set smaller than the formula's letters, δ at their size
    found = scripts.kinds([_line(("∑", 0, 10, 14.8, 118), ("k", 10, 13, 5.9, 124), ("δ", 15, 21, 10.3, 116),
                                 ("x", 22, 27, 10.3, 116), main=(116, 10.3))])
    assert found == [["", "sub", "", ""]]


def test_a_glyph_across_a_column_gap_is_no_script():
    # the left column's last ideograph (its box wider than its step), a full stop inside that box, and the right
    # column's smaller text 3 pt higher, 100 pt away
    left = _line(("值", 162, 183, 10.5, 581.4), main=(581.4, 10.5))
    stop = _line(("．", 172, 175, 10.3, 581.4), main=(581.4, 10.3))
    right = _line(*_row("test", 280, 8.8, 578.1, 3), main=(578.1, 8.8))
    assert scripts.kinds([left, stop, right])[2] == ["", "", "", ""]


def test_a_script_before_its_base_is_one_before_a_letter_or_digit_only():
    # ¹³C and a footnote mark before a word are scripts of what follows; a small digit before a closing parenthesis
    # whose own base is out of the row is not judged
    nuclide = scripts.kinds([_line(("1", 0, 3, 6, 96), ("3", 3, 6, 6, 96), ("C", 6, 13, 10, 100), main=(100, 10))])
    before_paren = scripts.kinds([_line(("2", 0, 3, 6, 103), (")", 3, 7, 10, 100), ("+", 8, 14, 10, 100),
                                        ("x", 15, 21, 10, 100), main=(100, 10))])
    assert nuclide == [["sup", "sup", ""]]
    assert before_paren == [["", "", "", ""]]


def test_unreadable_glyphs_alone_are_no_script_run():
    s = scripts.SUP_MARK
    assert scripts.runs(scripts.unmark(f"方法 B{s}\ue000")) == []
    assert scripts.write(f"[{s}1{s}\ue000{s}3]") == "[$^{1\ue0003}$]"


def test_no_script_of_a_glyph_across_a_line_of_another_row():
    heading, body = _line(*_row("论", 45, 27, 98.5)), _line(*_row("将铰缝刚度均作为", 72, 20, 118))
    reference = _line(*_row("[3]李海生", 256, 18, 106.8))
    assert scripts.kinds([heading, body, reference])[2] == [""] * 6


# ── writing ──


def test_a_run_is_written_in_unicode_where_every_character_has_a_form_else_in_latex():
    assert scripts.form("sub", "1") == "₁"
    assert scripts.form("sup", "−1") == "⁻¹"
    assert scripts.form("sub", "sd") == "$_{sd}$"
    assert scripts.form("sup", "[1, 3]") == "$^{[1, 3]}$"
    assert scripts.form("sup", "*") == "$^{*}$"
    assert scripts.form("sub", "升") == "$_{\\text{升}}$"


def test_marked_text_is_written_with_the_gap_before_a_script_removed():
    s, p = scripts.SUB_MARK, scripts.SUP_MARK
    assert scripts.write(f"k{s}1 与 m {p}2") == "k₁ 与 m²"
    assert scripts.write(f"梁{p}[{p}1{p}] 等") == "梁$^{[1]}$ 等"
    assert scripts.write("纯文字 k1") == "纯文字 k1"


def test_a_script_mark_is_found_after_the_glyph_before_it():
    marks = [Mark(kind="sub", text="1", before="k"), Mark(kind="sup", text="2", before="m")]
    assert emphasize("表1中 k1 与 m 2", marks)[0] == "表1中 k₁ 与 m²"
    assert emphasize("梁[1]等", [Mark(kind="sup", text="[1]", before="梁")])[0] == "梁$^{[1]}$等"


def test_a_script_inside_bold_is_written_inside_it():
    marks = [Mark(kind="bold", text="CO2 排放"), Mark(kind="sub", text="2", before="O")]
    assert emphasize("CO2 排放与其他", marks)[0] == "**CO₂ 排放**与其他"


def test_a_mark_not_found_leaves_the_text_unchanged():
    text, left = emphasize("$k_{1}$ 的取值", [Mark(kind="sub", text="1", before="k")])
    assert text == "$k_{1}$ 的取值" and len(left) == 1


# ── extraction ──


def _scripted_pdf(path):
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 100), "area 12 m", fontsize=11)
    page.insert_text((72 + pymupdf.get_text_length("area 12 m", fontsize=11), 96), "2", fontsize=7)
    page.insert_text((72, 140), "the value of k", fontsize=11)
    page.insert_text((72 + pymupdf.get_text_length("the value of k", fontsize=11), 143), "1", fontsize=7)
    doc.save(path)


def test_scripts_of_native_text_are_marks_after_their_base(tmp_path):
    _scripted_pdf(tmp_path / "s.pdf")
    blocks = [b for b in extract_pdf(tmp_path / "s.pdf").blocks if b.kind == BlockKind.TEXT]
    written = [emphasize(b.text, b.observations[0].marks)[0] for b in blocks]
    assert "".join(b.text for b in blocks).replace(" ", "").count("m2") == 1  # the text itself stays as read
    assert any("12 m²" in w for w in written) and any("k₁" in w for w in written)


def test_scripts_in_a_ruled_table_are_written_into_the_cells(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    xs, ys = [72, 200, 328], [100, 130, 160]
    for x in xs:
        page.draw_line((x, ys[0]), (x, ys[-1]))
    for y in ys:
        page.draw_line((xs[0], y), (xs[-1], y))
    page.insert_text((77, 120), "k", fontsize=11)
    page.insert_text((77 + pymupdf.get_text_length("k", fontsize=11), 123), "1", fontsize=7)
    page.insert_text((205, 120), "k2", fontsize=11)
    for c, text in enumerate(("0.5", "10")):
        page.insert_text((xs[c] + 5, 150), text, fontsize=11)
    doc.save(tmp_path / "t.pdf")
    table = next(b for b in extract_pdf(tmp_path / "t.pdf").blocks if b.kind == BlockKind.TABLE)
    assert [[c.content for c in row] for row in table.cells.slot_matrix()] == [["k₁", "k2"], ["0.5", "10"]]


def test_no_scripts_from_geometry_inside_a_display_formula(tmp_path):
    _scripted_pdf(tmp_path / "s.pdf")
    formula = ("display_formula", (60.0, 125.0, 300.0, 150.0))  # around "the value of k₁"
    ext = extract_pdf(tmp_path / "s.pdf", layout=lambda page: [formula])
    marks = [(m.kind, m.text) for b in ext.blocks for o in b.observations for m in o.marks]
    assert ("sup", "2") in marks and ("sub", "1") not in marks


def test_word_runs_set_as_scripts_are_marks(tmp_path):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE

    from parserx.content.docx import extract_docx

    doc = Document()
    raised = doc.styles.add_style("Raised", WD_STYLE_TYPE.CHARACTER)
    raised.font.superscript = True
    para = doc.add_paragraph("面积 12 m")
    para.add_run("2").font.superscript = True
    para.add_run("，水 H")
    para.add_run("2").font.subscript = True
    para.add_run("O，注")
    para.add_run("1").style = raised
    doc.save(tmp_path / "s.docx")
    block = extract_docx(tmp_path / "s.docx").blocks[0]
    marks = block.observations[0].marks
    assert [(m.kind, m.text, m.before) for m in marks] == [("sup", "2", "m"), ("sub", "2", "H"), ("sup", "1", "注")]
    assert emphasize(block.text, marks)[0] == "面积 12 m²，水 H₂O，注¹"


def test_a_joined_paragraph_keeps_the_scripts_of_each_part():
    # milestone audit (2026-10-02): paper_chn02's "3.886 × 10⁻⁴ m²/kN" ran on into the next column; after the agent
    # joined the two blocks, the second's scripts were lost — the joined paragraph rendered with the first's marks only
    from parserx.ir.anchor import PdfAnchor
    from parserx.ir.block import Block
    from parserx.ir.enums import BlockKind, DocumentStatus, ObservationStatus, PageStatus, RelationKind, TaskKind
    from parserx.ir.observation import Observation
    from parserx.ir.relation import Relation
    from parserx.ir.state import DocumentState, PageState
    from parserx.render import render_markdown

    def block(bid, order, text, marks):
        anchor = PdfAnchor(page=1, bbox=(0, order * 20, 400, order * 20 + 12), coord_space="page_pt")
        obs = Observation(id=f"o-{bid}", engine="native_pdf", engine_version="v", task=TaskKind.EXTRACT, anchor=anchor,
                          text=text, marks=marks, status=ObservationStatus.OK)
        return Block(id=bid, kind=BlockKind.TEXT, order=order, anchors=[anchor], text=text, observations=[obs],
                     chosen_observation=obs.id)

    first = block("a", 0, "板的理论参数 wi =3.886 ×", [Mark(kind="sub", text="i", before="w")])
    second = block("b", 1, "10-4 m2 /kN，φi =1.574", [Mark(kind="sup", text="-4", before="10"),
                                                      Mark(kind="sup", text="2", before="m"),
                                                      Mark(kind="sub", text="i", before="φ")])
    state = DocumentState(id="d", source="d.pdf", source_sha256="0" * 64, format="pdf",
                          status=DocumentStatus.IN_PROGRESS, pages=[PageState(n=1, unit="pdf_page", status=PageStatus.DONE)],
                          blocks=[first, second], relations=[Relation(id="r-1", kind=RelationKind.CONTINUES,
                                                                      src="a", dst="b")])
    text = render_markdown(state)
    assert not any(plain in text for plain in ("10-4", "m2", "φi", "wi")), text
