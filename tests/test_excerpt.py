"""Word excerpts for test data (eval/excerpt.py): chosen body elements kept, unused images removed."""

import io
import zipfile

import docx
import pytest
from PIL import Image

from parserx.eval.excerpt import excerpt_docx, parse_parts


def _png(colour):
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), colour).save(buf, "PNG")
    buf.seek(0)
    return buf


def _source(path):
    d = docx.Document()
    d.add_paragraph("zero")
    d.add_paragraph().add_run().add_picture(_png("red"))
    d.add_paragraph("two")
    p = d.add_paragraph()
    for colour in ("green", "blue", "white"):
        p.add_run().add_picture(_png(colour))
    d.add_paragraph("four")
    d.save(path)


def test_parts_are_ranges_single_elements_or_an_element_with_its_first_images():
    assert [(p.first, p.last, p.images) for p in parse_parts(["0-2", "5", "7:3"])] == [(0, 2, None), (5, 5, None), (7, 7, 3)]
    with pytest.raises(ValueError):
        parse_parts(["4-6", "5"])
    with pytest.raises(ValueError):
        parse_parts(["3-1"])


def test_an_excerpt_keeps_the_chosen_elements_and_only_the_images_they_show(tmp_path):
    src, dst = tmp_path / "a.docx", tmp_path / "b.docx"
    _source(src)
    counts = excerpt_docx(src, dst, parse_parts(["0", "2", "3:2"]))
    kept = docx.Document(dst)
    assert [p.text for p in kept.paragraphs] == ["zero", "two", ""]
    assert counts["images"] == 2
    media = [n for n in zipfile.ZipFile(dst).namelist() if n.startswith("word/media/")]
    assert len(media) == 2  # the red image and the third of the row are gone
    assert len(kept.inline_shapes) == 2
