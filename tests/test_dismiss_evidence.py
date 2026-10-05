"""What a look counts as evidence of (agent runs 2026-10-05, tr3): the agent's looks were refused where they did show
the place — a block read inside the image an item names, two looks cited together, a page of the document for an
item on its outline — and an item on a Word document's own text could not be closed at all: it has no image."""

from parserx.ir import ids
from parserx.ir.anchor import AssetAnchor, DocxAnchor, PdfAnchor
from parserx.ir.block import Block
from parserx.ir.enums import BlockKind, DocumentStatus, RelationKind
from parserx.ir.evidence import Evidence
from parserx.ir.relation import Relation
from parserx.ir.state import DocumentState
from parserx.tools.evidence import cited, document_evidence, image_evidence, word_text
from parserx.tools import workspace_init
from parserx.workspace.store import Workspace
from tests.test_tools_contract import _call, _config, _context, _edit


def _state():
    figure = Block(id="b-d00001", kind=BlockKind.FIGURE, order=0,
                   anchors=[DocxAnchor(part="word/document.xml", node_path="/w:body/w:p[1]")])
    inside = Block(id="b-d00001-r001", kind=BlockKind.TEXT, order=1, text="2025年度",
                   anchors=[AssetAnchor(asset="a-1", bbox=(0, 0, 10, 10), image_size=(100, 100))])
    page = Block(id="b-p001-0001", kind=BlockKind.TEXT, order=2, text="正文",
                 anchors=[PdfAnchor(page=1, bbox=(0, 0, 100, 12), coord_space="page_pt")])
    return DocumentState(
        id="d", source="d.docx", source_sha256="0" * 64, format="docx", status=DocumentStatus.IN_PROGRESS,
        blocks=[figure, inside, page],
        relations=[Relation(id=ids.relation_id(RelationKind.CONTAINS, figure.id, inside.id),
                            kind=RelationKind.CONTAINS, src=figure.id, dst=inside.id)],
        evidence=[Evidence(id="e-inside", how="image", block=inside.id, image="a-crop"),
                  Evidence(id="e-page2", how="image", page=2, image="a-page2")])


def test_a_look_at_a_block_read_inside_an_image_shows_the_image():
    state = _state()
    assert image_evidence(state, state.blocks[0], "e-inside").passed
    assert not image_evidence(state, state.blocks[2], "e-inside").passed  # it shows nothing of another page


def test_looks_cited_together_are_each_evidence():
    state = _state()
    assert [e.id for e in cited(state, "e-unknown, e-inside，a-page2")] == ["e-inside", "e-page2"]
    assert image_evidence(state, state.blocks[0], "e-page2,e-inside").passed


def test_an_item_on_the_whole_document_rests_on_any_look_at_it():
    state = _state()
    assert document_evidence(state, "e-page2").passed
    assert not document_evidence(state, "").passed


def test_a_word_documents_own_text_is_the_source_and_its_images_are_not():
    state = _state()
    assert [word_text(b) for b in state.blocks] == [False, False, False]  # an image, read inside one, a PDF block
    paragraph = Block(id="b-d00002", kind=BlockKind.TEXT, order=3, text="（一）企业情况简介",
                      anchors=[DocxAnchor(part="word/document.xml", node_path="/w:body/w:p[2]")])
    assert word_text(paragraph)


def test_the_outline_item_of_a_word_document_is_closed_without_a_look(tmp_path):
    import docx

    document = docx.Document()
    document.add_heading("一、企业基本情况", level=1)
    document.add_paragraph("（一）企业情况简介。")
    document.add_paragraph("企业成立于二〇一〇年，主营桥梁监测设备的研发、生产与销售，拥有多项专利。")
    document.add_paragraph("（二）企业营业执照。")
    path = tmp_path / "w.docx"
    document.save(path)
    ws = tmp_path / "ws"
    workspace_init(path, ws, config=_config())
    context = _context()
    issues = _call("read_draft", ws, {"view": "issues"}, context=context)[0].result.issues
    item = next(u for u in issues if u.kind == "outline_review")
    assert word_text(next(b for b in Workspace.open(ws).load().blocks if b.id == item.target))
    outcome = _edit(ws, context, {"op": "dismiss", "issue": item.id, "reason": "条目是列表项，不是标题"})[0]
    assert outcome.accepted, outcome.detail
