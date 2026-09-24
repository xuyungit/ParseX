"""Vector images (EMF / WMF) rendered to PNG (P2-4 C4) on synthetic drawings."""

import io
import shutil
import struct
import zipfile

import fitz
import pytest
from docx import Document
from PIL import Image

from parserx.content import vector
from parserx.content.docx import extract_docx
from parserx.ir.anchor import AssetAnchor
from parserx.ir.enums import BlockKind

needs_libreoffice = pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")


def _wmf(w=2880, h=1440, inch=1440, color=(20, 90, 200)) -> bytes:
    """A placeable WMF of w × h units (inch units per inch) filled with one rectangle."""
    def rec(fn, *params):
        return struct.pack("<IH", 3 + len(params), fn) + b"".join(struct.pack("<h", p) for p in params)

    brush = struct.pack("<IHH4BH", 7, 0x02FC, 0, *color, 0, 0)
    body = (rec(0x0103, 8) + rec(0x020B, 0, 0) + rec(0x020C, h, w) + brush + rec(0x012D, 0)
            + rec(0x041B, h, w, 0, 0) + rec(0x0000))
    header = struct.pack("<HHHIHIH", 1, 9, 0x0300, 9 + len(body) // 2, 1, 7, 0)
    first = struct.pack("<IHhhhhHI", 0x9AC6CDD7, 0, 0, 0, w, h, inch, 0)
    checksum = 0
    for (word,) in struct.iter_unpack("<H", first):
        checksum ^= word
    return first + struct.pack("<H", checksum) + header + body


def _about(size, expected) -> bool:
    return all(abs(a - b) <= 1 for a, b in zip(size, expected))  # the clip is rounded outwards to whole pixels


def _pdf(draw) -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.draw_rect(page.rect, color=None, fill=(1, 1, 1))  # the page background LibreOffice paints
    draw(page)
    return doc.tobytes()


def test_the_drawing_is_rendered_without_the_page_around_it():
    png = vector.render_drawing(_pdf(lambda p: p.draw_rect(fitz.Rect(100, 200, 244, 272), color=None, fill=(0, 0, 1))))
    with Image.open(io.BytesIO(png)) as image:
        assert _about(image.size, (400, 200))  # 2 × 1 inch at 200 dpi
        assert image.getpixel((200, 100)) == (0, 0, 255)


def test_a_blank_drawing_gives_nothing():
    assert vector.render_drawing(_pdf(lambda p: None)) is None


def test_large_drawings_are_capped(monkeypatch):
    monkeypatch.setattr(vector, "MAX_SIDE", 1000)
    png = vector.render_drawing(_pdf(lambda p: p.draw_rect(fitz.Rect(0, 0, 595, 421), color=None, fill=(1, 0, 0))))
    with Image.open(io.BytesIO(png)) as image:
        assert _about(image.size, (1000, 708))


@needs_libreoffice
def test_vector_images_render_to_png_deterministically():
    items = {"a": (_wmf(), "image/x-wmf"), "b": (b"not a drawing", "image/x-emf")}
    rendered = vector.render_vectors(items)
    assert set(rendered.images) == {"a"} and rendered.version  # bytes LibreOffice opens as text are left out
    with Image.open(io.BytesIO(rendered.images["a"])) as image:
        assert _about(image.size, (400, 200)) and image.getpixel((200, 100)) == (20, 90, 200)
    assert vector.render_vectors(items) == rendered


def _docx_with_wmf(tmp_path):
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), (255, 255, 255)).save(buf, "PNG")
    doc = Document()
    doc.add_paragraph("Before")
    doc.add_picture(io.BytesIO(buf.getvalue()))
    source, path = tmp_path / "png.docx", tmp_path / "wmf.docx"
    doc.save(source)
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(path, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            name = item.filename
            if name.startswith("word/media/"):
                name, data = name.rsplit(".", 1)[0] + ".wmf", _wmf()
            elif name in ("word/_rels/document.xml.rels", "[Content_Types].xml"):
                data = data.replace(b".png", b".wmf").replace(b'Extension="png" ContentType="image/png"',
                                                              b'Extension="wmf" ContentType="image/x-wmf"')
            zout.writestr(name, data)
    return path


def _figure_asset(ext):
    figure = next(b for b in ext.blocks if b.kind == BlockKind.FIGURE)
    anchor = next(a for a in figure.anchors if isinstance(a, AssetAnchor))
    return figure, next(a for a in ext.assets if a.id == anchor.asset)


@needs_libreoffice
def test_docx_vector_images_become_png_assets(tmp_path):
    ext = extract_docx(_docx_with_wmf(tmp_path))
    figure, asset = _figure_asset(ext)
    assert asset.media_type == "image/png" and _about((asset.width, asset.height), (400, 200))
    assert figure.decisions[0].evidence == {"rendered_from": "image/x-wmf", "renderer": "libreoffice"}
    assert ext.engines["libreoffice"] and not ext.warnings


def test_without_libreoffice_the_original_is_kept_with_a_warning(tmp_path, monkeypatch):
    monkeypatch.setattr(vector.shutil, "which", lambda name: None)
    ext = extract_docx(_docx_with_wmf(tmp_path))
    _figure, asset = _figure_asset(ext)
    assert asset.media_type == "image/x-wmf"
    assert any("EMF / WMF" in w for w in ext.warnings)


def test_libreoffice_is_given_the_system_fonts(tmp_path, monkeypatch):
    # its own fontconfig may have no configuration (macOS build): only its bundled fonts, no Chinese
    monkeypatch.delenv("FONTCONFIG_FILE", raising=False)
    env = vector.soffice_env(tmp_path)
    conf = (tmp_path / "fonts.conf").read_text()
    assert env["FONTCONFIG_FILE"] == str(tmp_path / "fonts.conf")
    assert "/etc/fonts/fonts.conf" in conf and "<dir>/System/Library/Fonts</dir>" in conf
    assert "<dir>/usr/share/fonts</dir>" in conf
    monkeypatch.setenv("FONTCONFIG_FILE", "/somewhere/fonts.conf")
    assert vector.soffice_env(tmp_path / "other")["FONTCONFIG_FILE"] == "/somewhere/fonts.conf"

