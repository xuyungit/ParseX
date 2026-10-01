"""A TOC entry's inherited heading outline does not make it a heading."""

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from parserx.content.docx import extract_docx
from parserx.hierarchy.docx_styles import propose_docx_structure


def test_toc_entries_do_not_inherit_heading_roles(tmp_path):
    doc = Document()
    doc.add_paragraph("Report", style="Title")
    doc.add_paragraph("Contents", style="Heading 1")
    for name in ("toc 1", "TOC2", "目录 1"):
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Heading 1"]
        doc.add_paragraph("Scope ........ 3", style=name)
    doc.add_paragraph("Scope", style="Heading 1")
    path = tmp_path / "toc.docx"
    doc.save(path)
    state = extract_docx(path).to_state(doc_id="d", source="toc.docx", source_sha256="0" * 64)
    text = {b.id: b.text for b in state.blocks}
    roles = {text[c["block"]]: c["role"] for c in propose_docx_structure(state)}
    assert roles == {"Report": "H1", "Contents": "H2", "Scope": "H2"}
