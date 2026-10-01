"""Tool contracts (guide §5, interfaces §4–§5): envelope, DocText, failure codes, versions, legality, CLI."""

import io
import json
import sys
from pathlib import Path

import pymupdf
import pytest
from jsonschema import Draft202012Validator
from PIL import Image

import parserx.cli
from parserx.config.schema import CacheConfig, OCRBuilderConfig, ParserXConfig, PriceConfig
from parserx.ir.enums import BlockKind, PageStatus
from parserx.render import render_markdown
from parserx.services.ocr import PaddleOCRService
from parserx.tools import TOOLS, ToolContext, agent_json, call_tool, tool_schema, workspace_init
from parserx.workspace import Workspace, verify_workspace

NATIVE = "SENTINEL-NATIVE 采购金额为 100 万元"
OCR_TEXT = "SENTINEL-OCR 扫描文字 3 件"
VLM_TEXT = "SENTINEL-VLM 一座桥"
SENTINELS = ("SENTINEL-NATIVE", "SENTINEL-OCR", "SENTINEL-VLM")


def _png(w, h, color):
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, "PNG")
    return buf.getvalue()


@pytest.fixture
def pdf(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 90), "SENTINEL-NATIVE 标题", fontsize=18, fontname="china-s")
    page.insert_text((72, 140), NATIVE, fontsize=11, fontname="china-s")
    page.insert_image(pymupdf.Rect(72, 300, 200, 400), stream=_png(128, 100, (200, 40, 40)))
    scanned = doc.new_page(width=595, height=842)
    scanned.insert_image(scanned.rect, stream=_png(300, 424, (230, 230, 230)))
    path = tmp_path / "doc.pdf"
    doc.save(path)
    return path


def _ocr_page():
    return {"prunedResult": {"width": 1000, "height": 1400, "parsing_res_list": [
        {"block_label": "paragraph_title", "block_content": "SENTINEL-OCR 标题", "block_bbox": [100, 100, 600, 140],
         "block_order": 1},
        {"block_label": "text", "block_content": OCR_TEXT, "block_bbox": [100, 160, 900, 220], "block_order": 2},
        {"block_label": "table", "block_content":
            "<table><tr><td>项目</td><td>数值</td></tr><tr><td>SENTINEL-OCR 甲</td><td>3</td></tr></table>",
         "block_bbox": [100, 300, 900, 500], "block_order": None},
        {"block_label": "image", "block_content": "", "block_bbox": [100, 600, 500, 900], "block_order": None},
        {"block_label": "number", "block_content": "2", "block_bbox": [480, 1350, 520, 1380], "block_order": None},
    ]}}


class FakeVLM:
    def __init__(self):
        self.attempt_hook = None
        self.usage_hook = None
        self.calls = []
        self.schemas = []
        self.figure_type = "photo"  # what a description says the image is (IO6: content is transcribed)
        self.caption = VLM_TEXT

    def describe_image(self, image_path, prompt, *, context="", temperature=0.1, max_tokens=8192,
                       structured_output_mode="off", json_schema=None, json_schema_name="x"):
        self.calls.append(json_schema_name)
        self.schemas.append(json_schema)
        image_path.read_bytes()  # like the real service: the image file is read (a str path would fail here)
        if self.usage_hook:
            self.usage_hook("gpt-6-luna", 1000, 0, 100)
        if json_schema_name == "parserx_ask_image":
            return "SENTINEL-VLM 图上写的是：扫描文字 8 件"
        if json_schema_name == "parserx_review_table":
            return json.dumps({"table_html": "<table><tr><td>项目</td><td>数值</td></tr>"
                                             "<tr><td>SENTINEL-OCR 甲</td><td>8</td></tr></table>",
                               "undetermined": []})
        return json.dumps({"type": self.figure_type, "caption": self.caption})


class FakeReader:
    """The local page reader without the model: it sees nothing unless told (tests/test_page_reading.py)."""

    name = "reading"
    version = "fake-reader-1"

    def __init__(self, lines=None):
        self.lines, self.calls = lines or {}, 0

    def read(self, png):
        self.calls += 1
        return list(self.lines.get(self.calls, []))


def _context(ocr_behaviour=None, page=None):
    vlm = FakeVLM()

    class Context(ToolContext):
        def _new_reader(self):
            return FakeReader()

        def _new_ocr(self):
            service = PaddleOCRService(OCRBuilderConfig(endpoint="https://x/api/v2/ocr/jobs", token="t"))

            def run_job(file_bytes, filename, mime, job_key=None):
                if ocr_behaviour is not None:
                    raise ocr_behaviour
                with pymupdf.open(stream=file_bytes, filetype="pdf") as sub:
                    return {"layoutParsingResults": [(page or _ocr_page)() for _ in sub]}

            service._run_job = run_job
            return service

        def _new_vlm(self, cfg):
            return vlm

    Context.fake_vlm = vlm
    return Context


def _config(**budget):
    config = ParserXConfig()
    config.services.vlm.model = "gpt-6-luna"
    config.scheduling.retry.backoff_s = 0.0
    config.scheduling.budget.requests = budget
    config.scheduling.prices = {"gpt-6-luna": PriceConfig(input=0.10, cached_input=0.01, output=0.50)}
    config.output.report = config.output.sidecar = True  # the tests read the whole package
    return config


@pytest.fixture
def ws(pdf, tmp_path):
    envelope, code = workspace_init(pdf, tmp_path / "ws", config=_config())
    assert code == 0 and envelope.ok
    return tmp_path / "ws"


def _call(name, ws, request=None, *, config=None, context=None, expect_version=None):
    return call_tool(name, ws, request or {}, config=config or _config(), context_factory=context or _context(),
                     expect_version=expect_version)


def _strings_outside_doc_text(value, inside=False):
    if isinstance(value, dict):
        for key, item in value.items():
            yield from _strings_outside_doc_text(item, inside or key == "doc_text")
    elif isinstance(value, list):
        for item in value:
            yield from _strings_outside_doc_text(item, inside)
    elif isinstance(value, str) and not inside:
        yield value


def _assert_contract(envelope, name):
    data = json.loads(agent_json(envelope))  # what the agent reads
    Draft202012Validator(tool_schema(name)["envelope"]).validate(data)
    leaked = [s for s in _strings_outside_doc_text(data) if any(x in s for x in SENTINELS)]
    assert leaked == [], f"document text outside DocText in {name}: {leaked[:3]}"
    return data


# ── Helpers for the agent's tools ───────────────────────────────────────


def _look(ws, context, **look):
    """One look at the source; returns its result (with the evidence id)."""
    env, _ = _call("view_source", ws, {"looks": [look]}, context=context)
    assert env.ok, env.failures
    return env.result.results[0]


def _evidence(ws, context, **look):
    return _look(ws, context, **look).evidence


def _edit(ws, context, *ops, atomic=False):
    env, code = _call("edit_draft", ws, {"ops": list(ops), "atomic": atomic}, context=context)
    assert env.ok and code == 0, env.failures
    return env.result.outcomes


def _issues(ws, context):
    return _call("read_draft", ws, {"view": "issues"}, context=context)[0].result.issues


def _accounts(ws):
    from parserx.accounting import check

    return check(Workspace.open(ws).load(), ws)


# ── A full session through the agent's tools ────────────────────────────


def test_session_through_every_tool(ws, tmp_path):
    context = _context()
    env, _ = _call("read_draft", ws, context=context)
    data = _assert_contract(env, "read_draft")
    assert [p["status"] for p in data["result"]["summary"]["pages"]] == ["done", "pending"]
    assert data["result"]["summary"]["issues"] == {"page_pending": 1}

    env, _ = _call("view_source", ws, {"looks": [{"page": 1, "as": "image"}]}, context=context)
    data = _assert_contract(env, "view_source")
    look = data["result"]["results"][0]
    assert look["image"]["path"].endswith(".png") and look["evidence"].startswith("e-")
    assert data["ws_version"] == 2  # the look is kept as evidence; the draft is unchanged

    env, _ = _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    data = _assert_contract(env, "recognize")
    assert env.ok and data["result"]["selections"][0]["choice"] == "scan_engine"
    assert data["cost"]["requests"] == {"ocr": 1} and Workspace.open(ws).load().pages[1].status == "done"

    table = next(b for b in Workspace.open(ws).load().blocks if b.kind == BlockKind.TABLE)
    reading = _look(ws, context, block=table.id, **{"as": "table"},
                    issues=[{"kind": "char", "cells": [[1, 1]], "note": "3 or 8?"}])
    figure = next(b for b in Workspace.open(ws).load().blocks if b.kind == BlockKind.FIGURE)
    description = _look(ws, context, block=figure.id, **{"as": "description"})
    title = next(b for b in Workspace.open(ws).load().blocks if b.kind == BlockKind.TITLE)
    env, _ = _call("edit_draft", ws, {"ops": [
        {"op": "adopt", "block": table.id, "evidence": reading.evidence, "reason": "图上是 8"},
        {"op": "adopt", "block": figure.id, "evidence": description.evidence, "reason": "描述"},
        {"op": "set_role", "block": title.id, "role": "H1", "reason": "engine title"},
        {"op": "set_role", "block": table.id, "role": "H2", "reason": "illegal"}]}, context=context)
    data = _assert_contract(env, "edit_draft")
    assert [o["accepted"] for o in data["result"]["outcomes"]] == [True, True, True, False]
    assert data["result"]["outcomes"][3]["rule"] == "kind_not_structural"

    env, _ = _call("submit_draft", ws, context=context)
    data = _assert_contract(env, "submit_draft")
    assert data["result"]["accepted"] and data["result"]["status"] == "complete"
    env, _ = _call("export", ws, {"out": str(tmp_path / "out")}, context=context)
    assert _assert_contract(env, "export")["result"]["accepted"]
    markdown = (tmp_path / "out" / "doc.md").read_text()
    assert "# SENTINEL-OCR 标题" in markdown and f"![照片](" in markdown and f"> 图片说明：{VLM_TEXT}" in markdown and "| SENTINEL-OCR 甲 | 8 |" in markdown

    state = Workspace.open(ws).load()
    assert state.stats.requests == {"ocr": 1, "vlm": 2} and state.stats.cost_usd == pytest.approx(2 * (1000 * 0.10 + 100 * 0.50) / 1e6)
    calls = [json.loads(line) for line in (ws / "calls.jsonl").read_text().splitlines()]
    assert [c["tool"] for c in calls if c["type"] == "call"][:3] == ["workspace_init", "read_draft", "view_source"]
    assert verify_workspace(ws).ok  # every commit came from a tool call


def test_a_workspace_changed_outside_the_tools_is_refused(ws, tmp_path):
    raw = json.loads((ws / "state.json").read_text())
    raw["blocks"][0]["text"] = "SENTINEL-NATIVE 改写"
    (ws / "state.json").write_text(json.dumps(raw, ensure_ascii=False))
    for name, request in (("submit_draft", {}), ("export", {"out": str(tmp_path / "out")}), ("read_draft", {})):
        env, code = _call(name, ws, request)
        assert not env.ok and code == 0 and env.failures[0].code == "workspace_tampered", name
        assert not env.failures[0].retryable
    assert not (tmp_path / "out").exists()


def test_schemas_exist_and_structure_changes_carry_no_text():
    from pydantic import TypeAdapter

    from parserx.hierarchy import StructureChange

    for name in TOOLS:
        schema = tool_schema(name)
        assert schema["request"]["type"] == "object" and "properties" in schema["envelope"]
    changes = TypeAdapter(list[StructureChange]).json_schema()["$defs"]
    assert changes and not any("text" in d.get("properties", {}) for d in changes.values())


# ── Failure codes ───────────────────────────────────────────────────────


def test_invalid_request_exits_two(ws):
    env, code = _call("read_draft", ws, {"view": "blocks"})
    assert code == 2 and not env.ok and env.failures[0].code == "invalid_request"


def test_not_found(ws):
    env, code = _call("read_draft", ws, {"view": "blocks", "blocks": ["b-nope"]})
    assert code == 0 and not env.ok and env.failures[0].code == "not_found"


def test_version_conflict(ws):
    env, _ = _call("read_draft", ws, expect_version=99)
    assert not env.ok and env.failures[0].code == "version_conflict" and env.failures[0].retryable


def test_budget_exhausted_skips_the_page_but_exports_partial(ws, tmp_path):
    config = _config(ocr=0)
    env, _ = _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, config=config)
    assert env.ok and env.failures[0].code == "budget_exhausted"
    state = Workspace.open(ws).load()
    assert state.pages[1].status == PageStatus.SKIPPED
    env, _ = _call("export", ws, {"out": str(tmp_path / "out")}, config=config)
    assert env.ok and env.result.accepted and env.result.status == "partial"
    assert Workspace.open(ws).load().missing


@pytest.mark.parametrize("exc, code, retryable", [
    (RuntimeError("OCR submit HTTP 400: bad file"), "service_error", False),
    (TimeoutError("job still running"), "timeout", True),
])
def test_service_failures_mark_the_page(ws, exc, code, retryable):
    env, _ = _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=_context(exc))
    assert env.ok and (env.failures[0].code, env.failures[0].retryable) == (code, retryable)
    assert Workspace.open(ws).load().pages[1].status == PageStatus.FAILED


def test_offline_cache_miss_leaves_the_page_pending_and_submit_refused(ws, tmp_path):
    config = _config()
    config.cache = CacheConfig(mode="read_only", dir=str(tmp_path / "empty-cache"))
    env, _ = _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, config=config)
    assert env.failures[0].code == "cache_miss_offline"
    assert Workspace.open(ws).load().pages[1].status == PageStatus.PENDING
    env, _ = _call("export", ws, {"out": str(tmp_path / "out")}, config=config)
    assert env.ok and not env.result.accepted and "pending pages [2]" in env.result.blockers[0]
    assert not (tmp_path / "out").exists()


def test_internal_error_exits_one(ws, monkeypatch):
    spec = TOOLS["read_draft"]
    monkeypatch.setitem(TOOLS, "read_draft", spec.__class__(spec.request, spec.result, lambda ctx, req: 1 / 0))
    env, code = _call("read_draft", ws)
    assert code == 1 and env.failures[0].code == "internal_error"


def test_native_pages_are_not_sent_to_the_scan_engine(ws):
    env, _ = _call("recognize", ws, {"pages": [1], "engine": "paddleocr"})
    assert env.ok and env.cost.requests == {} and "native text layer passed" in env.failures[0].message


# ── CLI ─────────────────────────────────────────────────────────────────


def test_cli_prints_only_the_envelope(ws, monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["parserx", "dev", "tool", "read_draft", "--ws", str(ws), "--json"])
    with pytest.raises(SystemExit) as exit_info:
        parserx.cli.main()
    envelope = json.loads(capsys.readouterr().out)
    assert exit_info.value.code == 0 and envelope["tool"] == "read_draft" and envelope["ok"]


def test_cli_schema_and_invalid_request(ws, monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["parserx", "dev", "tool", "schema", "view_source"])
    with pytest.raises(SystemExit):
        parserx.cli.main()
    assert "request" in json.loads(capsys.readouterr().out)
    monkeypatch.setattr(sys, "argv", ["parserx", "dev", "tool", "read_draft", "--ws", str(ws), "--view", "blocks", "--json"])
    with pytest.raises(SystemExit) as exit_info:
        parserx.cli.main()
    assert exit_info.value.code == 2 and json.loads(capsys.readouterr().out)["failures"][0]["code"] == \
        "invalid_request"


def test_cli_looks_at_the_source_with_options(ws, monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["parserx", "dev", "tool", "view_source", "--ws", str(ws), "--page", "1", "--as",
                                      "image", "--json"])
    with pytest.raises(SystemExit) as exit_info:
        parserx.cli.main()
    envelope = json.loads(capsys.readouterr().out)
    assert exit_info.value.code == 0 and envelope["result"]["results"][0]["evidence"].startswith("e-")


def test_cli_options_are_the_request_fields():
    from parserx.tools.cli import _request

    def request(*argv):
        args = parserx.cli.build_parser().parse_args(["dev", "tool", *argv, "--ws", "ws"])
        return _request(args.tool_name, args)

    assert request("read_draft", "--view", "text", "--start", "b-p001-0002", "--after", "5", "--full") == \
        {"view": "text", "start": "b-p001-0002", "after": 5, "full": True}
    assert request("read_draft", "--view", "issues", "--kinds", "page_pending,title_candidate") == \
        {"view": "issues", "kinds": ["page_pending", "title_candidate"]}
    assert request("view_source", "--page", "2", "--bbox", "1", "2", "3", "4", "--as", "answer", "--question", "q") == \
        {"looks": [{"page": 2, "bbox": [1.0, 2.0, 3.0, 4.0], "as": "answer", "question": "q"}]}  # one look
    assert request("submit_draft") == {}


def test_empty_standard_input_is_named(ws, monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["parserx", "dev", "tool", "edit_draft", "--ws", str(ws), "--ops", "-", "--json"])
    monkeypatch.setattr(sys, "stdin", io.StringIO(""))
    with pytest.raises(SystemExit) as exit_info:
        parserx.cli.main()
    failure = json.loads(capsys.readouterr().out)["failures"][0]
    assert exit_info.value.code == 2 and "standard input is empty" in failure["message"]


# ── Reading the draft ───────────────────────────────────────────────────


def test_the_text_shows_what_the_output_shows(ws):
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"})
    lines = _call("read_draft", ws, {"view": "text", "page": 2})[0].result.lines
    state = Workspace.open(ws).load()
    hidden = {b.id for b in state.blocks if b.status in ("merged", "excluded", "duplicate")}
    assert lines and not {line.id for line in lines} & hidden
    merged = next(iter(b for b in state.blocks if b.status == "merged")).id
    detail = _call("read_draft", ws, {"view": "blocks", "blocks": [merged]})[0].result.blocks[0]
    assert detail.status == "merged"  # named blocks show whatever their status


def test_recognize_returns_views_only_when_asked(ws):
    env, _ = _call("recognize", ws, {"pages": [2], "engine": "paddleocr"})
    assert env.result.observations == [] and env.result.observations_total == 5


def test_answer_with_trailing_text_is_parsed():
    from parserx.tools.vlm_tasks import parse_review

    grid, undetermined, problem = parse_review(
        '{"table_html": "<table><tr><td>a</td></tr></table>", "undetermined": []}\n{"note": "extra"}')
    assert grid.slot(0, 0).content == "a" and problem is None


def test_a_rejected_table_reading_keeps_the_table(ws):
    class Rewriting(FakeVLM):
        def describe_image(self, *args, **kwargs):
            return json.dumps({"table_html": "<table><tr><td>项目</td></tr></table>", "undetermined": []})

    context = _context()
    rewriting = Rewriting()

    class Context(context):
        def _new_vlm(self, cfg):
            return rewriting

    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=Context)
    table = next(b for b in Workspace.open(ws).load().blocks if b.kind == BlockKind.TABLE)
    reading = _look(ws, Context, block=table.id, issues=[{"kind": "structure", "note": "rows?"}], **{"as": "table"})
    outcome = _edit(ws, Context, {"op": "adopt", "block": table.id, "evidence": reading.evidence, "reason": "r"})[0]
    assert not outcome.accepted and "structure_valid" in outcome.detail
    assert next(b for b in Workspace.open(ws).load().blocks if b.id == table.id).cells.n_rows == 2


def test_a_named_figure_schema_is_enforced(ws):
    context = _context()
    figure = next(b for b in Workspace.open(ws).load().blocks if b.kind == BlockKind.FIGURE)
    _call("describe_figure", ws, {"block": figure.id, "schema": "diagram"}, context=context)
    assert context.fake_vlm.schemas[-1]["properties"]["type"]["enum"] == ["diagram"]


# ── P2-5: batches and the standard processing ───────────────────────────


def test_describe_figures_in_one_batch(ws):
    context = _context()
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)  # page 2 brings a second figure
    figures = [b.id for b in Workspace.open(ws).load().blocks if b.kind == BlockKind.FIGURE and b.status == "ok"]
    text = next(b.id for b in Workspace.open(ws).load().blocks if b.kind == BlockKind.TEXT)
    assert len(figures) == 2
    env, code = _call("describe_figure", ws, {"blocks": [*figures, text]}, context=context)
    data = _assert_contract(env, "describe_figure")
    assert env.ok and code == 0 and data["cost"]["requests"] == {"vlm": 2}
    assert [i["block"] for i in data["result"]["items"]] == figures and all(i["type"] == "photo" for i in data["result"]["items"])
    assert [f["targets"] for f in data["failures"]] == [[text]] and data["failures"][0]["code"] == "invalid_request"
    state = Workspace.open(ws).load()
    assert all(b.semantic is not None for b in state.blocks if b.id in figures)
    env, _ = _call("describe_figure", ws, {"blocks": figures}, context=context)
    assert env.cost.requests == {} and all(i.cached for i in env.result.items)




def test_batch_results_do_not_depend_on_completion_order(ws, monkeypatch):
    import threading
    import time as _time

    context = _context()
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    figures = [b for b in Workspace.open(ws).load().blocks if b.kind == BlockKind.FIGURE and b.status == "ok"]
    fake = context.fake_vlm
    original = fake.describe_image
    first = threading.Event()

    def slow_first(image_path, prompt, **kw):
        # the first figure answers last; each answer names its own image
        if str(image_path).endswith(figures[0].anchors[-1].asset + ".png") or not first.is_set():
            first.set()
            _time.sleep(0.2)
        answer = json.loads(original(image_path, prompt, **kw))
        answer["caption"] = f"SENTINEL-VLM {Path(image_path).stem}"
        return json.dumps(answer)

    monkeypatch.setattr(fake, "describe_image", slow_first)
    env, _ = _call("describe_figure", ws, {"blocks": [f.id for f in figures]}, context=context)
    state = Workspace.open(ws).load()
    for block in (b for b in state.blocks if b.id in {f.id for f in figures}):
        asset = next(a.asset for a in block.anchors if hasattr(a, "asset"))
        assert block.semantic.caption.endswith(asset)


def test_batch_inputs_do_not_depend_on_other_tasks(pdf, tmp_path, monkeypatch):
    # §11.5 (v1's image-processor race): each request's image, prompt and context are fixed before any task
    # runs, so a batch whose first figure answers last sends exactly what describing each figure alone sends
    import threading
    import time as _time

    def requests(root, batches, slow_first):
        assert workspace_init(pdf, root, config=_config())[1] == 0
        context = _context()
        _call("recognize", root, {"pages": [2], "engine": "paddleocr"}, context=context)
        figures = sorted(b.id for b in Workspace.open(root).load().blocks if b.kind == BlockKind.FIGURE)
        fake, original, sent, first = context.fake_vlm, context.fake_vlm.describe_image, {}, threading.Event()

        def record(image_path, prompt, **kw):
            if slow_first and not first.is_set():
                first.set()
                _time.sleep(0.2)
            sent[Path(image_path).name] = (Path(image_path).read_bytes(), prompt, kw["context"])
            return original(image_path, prompt, **kw)

        monkeypatch.setattr(fake, "describe_image", record)
        for batch in batches(figures):
            _call("describe_figure", root, {"blocks": batch}, context=context)
        return sent

    together = requests(tmp_path / "a", lambda figures: [figures], slow_first=True)
    alone = requests(tmp_path / "b", lambda figures: [[f] for f in reversed(figures)], slow_first=False)
    assert len(together) >= 2 and together == alone



def test_the_pipeline_does_the_standard_steps_in_one_call(ws):
    config = _config()
    config.runtime.layout_shadow = False  # the layout step has its own tests (fake detector)
    context = _context()
    env, code = _call("run_pipeline", ws, {}, config=config, context=context)
    data = _assert_contract(env, "run_pipeline")
    result = data["result"]
    assert env.ok and code == 0 and data["cost"]["requests"] == {"ocr": 1, "vlm": 2}
    assert result["pages"] == {"done": 2} and result["figures"] == {"described": 2}
    assert result["check"]["exportable"] and result["check"]["document_status"] == "complete"
    assert [s["step"] for s in result["steps"]] == ["recognize", "reading", "describe_figure", "structure", "check"]
    calls = [json.loads(line) for line in (ws / "calls.jsonl").read_text().splitlines()]
    assert [c["tool"] for c in calls if c["type"] == "call"] == ["workspace_init", "run_pipeline"]
    assert verify_workspace(ws).ok
    again, _ = _call("run_pipeline", ws, {}, config=config, context=context)  # nothing left to do: no requests
    assert again.ok and again.cost.requests == {}


def test_the_pipeline_joins_a_paragraph_cut_by_the_page(tmp_path):
    doc = pymupdf.open()
    for text in ("供货方应在合同签订后分两批交货，第一批", "不少于总量的百分之六十。"):
        doc.new_page(width=595, height=842).insert_text((72, 400), text, fontsize=11, fontname="china-s")
    doc.save(tmp_path / "cut.pdf")
    config = _config()
    config.runtime.layout_shadow = False
    workspace_init(tmp_path / "cut.pdf", tmp_path / "cut", config=config)
    env, _ = _call("run_pipeline", tmp_path / "cut", {}, config=config)
    state = Workspace.open(tmp_path / "cut").load()
    assert [(r.kind.value, r.src, r.dst) for r in state.relations] == [("continues", "b-p001-0001", "b-p002-0001")]
    assert "1 paragraph continuations" in env.result.steps[-2].detail
    assert "第一批不少于" in render_markdown(state)


# ── Changing content: every change rests on evidence of its place ───────


def test_the_agent_corrects_ocr_text_it_has_seen(ws):
    context = _context()
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    block = next(b for b in Workspace.open(ws).load().blocks if b.text == OCR_TEXT)
    evidence = _evidence(ws, context, block=block.id)
    env, code = _call("edit_draft", ws, {"ops": [{"op": "replace_text", "block": block.id, "find": "3 件",
                                                  "replace": "8 件", "reason": "图上是 8 件", "evidence": evidence}]},
                      context=context)
    data = _assert_contract(env, "edit_draft")
    assert code == 0 and data["result"]["outcomes"][0]["accepted"] and data["cost"]["requests"] == {}
    after = next(b for b in Workspace.open(ws).load().blocks if b.id == block.id)
    assert after.text == OCR_TEXT.replace("3 件", "8 件") and after.chosen_observation.endswith("agent-1")
    assert after.decisions[-1].actor == "agent" and len(after.observations) == 2  # the OCR reading stays
    assert verify_workspace(ws).ok


def test_a_correction_needs_evidence_and_a_native_number_changed_is_recorded(ws):
    # execution plan §3.4: the agent may change a native number on the image's evidence; recorded, summarised
    from parserx.render.summary import document_summary

    context = _context()
    native = next(b for b in Workspace.open(ws).load().blocks if b.text == NATIVE)

    def replace(find, to, evidence):
        return _edit(ws, context, {"op": "replace_text", "block": native.id, "find": find, "replace": to,
                                   "reason": "图上是 900", "evidence": evidence})[0]

    outcome = replace("100 万元", "900 万元", "e-000000000000")
    assert not outcome.accepted and outcome.rule == "image_evidence"  # nothing was looked at: never
    evidence = _evidence(ws, context, block=native.id)
    outcome = replace("100 万元", "900 万元", evidence)
    assert outcome.accepted and "native numbers" in outcome.detail
    state = Workspace.open(ws).load()
    record = next(b for b in state.blocks if b.id == native.id).decisions[-1].evidence
    assert record["signal"] == "native_numbers_changed" and record["evidence"] == evidence
    assert "100" in record["signal_detail"] and "900" in record["signal_detail"]
    listed = document_summary(state, "doc").review.agent_overrides
    assert [(o.target, o.signals) for o in listed] == [(native.id, ["native_numbers_changed"])]
    assert "图上是 900" in listed[0].reason and listed[0].evidence == evidence
    assert replace("采购", "采买", evidence).accepted
    outcome = replace("不存在的字", "x", evidence)
    assert not outcome.accepted and outcome.rule == "find" and "exactly once" in outcome.detail


def _give_reading(ws, n, lines):
    """A local reading of page *n* (what the pipeline stores; tests give it directly)."""
    from parserx.ir.state import PageReading, ReadLine

    with Workspace.open(ws).txn("test:reading") as state:
        state.readings = sorted([r for r in state.readings if r.n != n] + [PageReading(
            n=n, engine="test", dpi=150, lines=[ReadLine(bbox=b, text=t, score=0.99) for t, b in lines])],
            key=lambda r: r.n)


def test_the_agent_adds_text_the_page_shows_and_no_block_has(ws):
    context = _context()
    native = next(b for b in Workspace.open(ws).load().blocks if b.text == NATIVE)
    _give_reading(ws, 1, [("SENTINEL-NATIVE 标题", (72, 72, 300, 95)), (NATIVE, native.anchors[0].bbox),
                          ("专家评审组名单", (72, 200, 200, 214))])

    def insert(text, bbox, evidence):
        return _edit(ws, context, {"op": "insert_text", "page": 1, "bbox": bbox, "text": text,
                                   "reason": "图上有这一行，文字层没有", "evidence": evidence})[0]

    assert not insert("专家评审组名单", [72, 200, 200, 214], "e-000000000000").accepted  # no image of the place
    page = _evidence(ws, context, page=1)
    outcome = insert("专家评审组名单", [72, 200, 200, 214], page)
    assert outcome.accepted and outcome.block
    state = Workspace.open(ws).load()
    added = next(b for b in state.blocks if b.id == outcome.block)
    assert added.text == "专家评审组名单" and added.kind == BlockKind.TEXT and added.anchors[0].bbox == (72, 200, 200, 214)
    order = [b.id for b in sorted(state.blocks, key=lambda b: b.order)]
    assert order.index(native.id) < order.index(added.id)  # placed by its position on the page
    entry = next(e for e in state.ledger if e.block == added.id)
    assert entry.unit == "agent_text" and entry.disposition == "output"
    accounts = _accounts(ws).accounting  # the reading was given outside a tool: check, not verify
    assert accounts.unassigned == 0 and accounts.output == accounts.discovered
    unseen = insert("图上没有的一行", [72, 600, 200, 614], page)  # §3.4: the local reading is a signal, not a gate
    assert unseen.accepted and "independent_reading" in unseen.detail
    added = next(b for b in Workspace.open(ws).load().blocks if b.id == unseen.block)
    assert added.decisions[-1].evidence["signal"] == "text_not_in_reading"
    assert not insert(NATIVE, list(native.anchors[0].bbox), page).accepted  # a block already has it
    assert sum(1 for b in Workspace.open(ws).load().blocks if b.text == NATIVE) == 1


def test_the_local_reading_of_a_native_number_is_told_not_obeyed(ws):
    """Where the local reading and the text layer agree on a number and nothing shows the agent's, the page prints it
    so: the change is refused (as printed).  A local reading that shows neither is told (a signal), not obeyed."""
    context = _context()
    native = next(b for b in Workspace.open(ws).load().blocks if b.text == NATIVE)
    edit = {"op": "replace_text", "block": native.id, "find": "100 万元", "replace": "900 万元", "reason": "图上是 900",
            "evidence": _evidence(ws, context, block=native.id)}
    _give_reading(ws, 1, [("SENTINEL-NATIVE 采购金额为100万元", native.anchors[0].bbox)])  # the page shows 100
    outcome = _edit(ws, context, edit)[0]
    assert not outcome.accepted and outcome.rule == "as_printed" and "agree with the draft" in outcome.detail
    _give_reading(ws, 1, [("SENTINEL-NATIVE 采购金额为700万元", native.anchors[0].bbox)])  # the reading has 700
    outcome = _edit(ws, context, {**edit, "evidence": _evidence(ws, context, block=native.id)})[0]
    assert outcome.accepted and "does not show" in outcome.detail
    detail = next(b for b in Workspace.open(ws).load().blocks if b.id == native.id).decisions[-1].evidence
    assert "does not show" in detail["signal_detail"]


# ── The worklist: an issue the evidence shows needs no change is dismissed ──


def test_an_issue_checked_on_the_source_can_be_dismissed(ws):
    context = _context()
    native = next(b for b in Workspace.open(ws).load().blocks if b.text == NATIVE)
    _give_reading(ws, 1, [(NATIVE, native.anchors[0].bbox), ("扫描软件的标志", (400.0, 780.0, 520.0, 792.0))])
    item = next(u for u in _issues(ws, context) if u.kind == "text_unaccounted")
    assert item.target == "p1" and item.id.startswith("w-")
    dismiss = {"op": "dismiss", "issue": item.id, "reason": "a logo the local reading took for text",
               "evidence": "e-000000000000"}
    outcome = _edit(ws, context, dismiss)[0]
    assert not outcome.accepted and outcome.rule == "image_evidence"  # the page was not looked at
    dismiss["evidence"] = _evidence(ws, context, page=1)
    env, code = _call("edit_draft", ws, {"ops": [dismiss]}, context=context)
    data = _assert_contract(env, "edit_draft")
    assert code == 0 and data["result"]["outcomes"][0]["accepted"] and data["result"]["issues_closed"] == [item.id]
    assert not any(u.kind == "text_unaccounted" for u in _issues(ws, context))
    pending = next(u for u in _issues(ws, context) if u.kind == "page_pending")
    outcome = _edit(ws, context, {**dismiss, "issue": pending.id})[0]
    assert not outcome.accepted and outcome.rule == "not_dismissable"  # resolved by processing, not dismissed
    outcome = _edit(ws, context, dismiss)[0]
    assert not outcome.accepted and outcome.rule == "unknown_issue"  # nothing open any more
    _give_reading(ws, 1, [(NATIVE, native.anchors[0].bbox), ("另一行漏掉的正文内容", (72.0, 600.0, 300.0, 612.0))])
    assert any(u.kind == "text_unaccounted" for u in _issues(ws, context))  # new content opens it again


def test_text_covered_on_the_page_stays_and_the_summary_names_it(ws):
    """Q71: text the page draws under another element is content; the agent keeps it and marks it occluded."""
    context = _context()
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    native = next(b for b in Workspace.open(ws).load().blocks if b.text == NATIVE)
    _give_reading(ws, 1, [("给助手发送消息", (72, 700, 300, 712))])  # the page image does not show NATIVE
    item = next(u for u in _issues(ws, context) if u.kind == "text_not_seen" and u.target == native.id)
    other = next(u for u in _issues(ws, context) if u.kind != "text_not_seen")
    page = _evidence(ws, context, page=1)
    outcome = _edit(ws, context, {"op": "dismiss", "issue": other.id, "reason": "r", "evidence": page,
                                  "occluded": True})[0]
    assert not outcome.accepted  # only text the page does not show can be occluded
    assert _edit(ws, context, {"op": "dismiss", "issue": item.id, "reason": "被悬浮的输入框盖住", "evidence": page,
                               "occluded": True})[0].accepted
    env, _ = _call("export", ws, {"out": str(ws.parent / "out"), "name": "d"}, context=context)
    assert env.ok and env.result.accepted, env.result
    assert NATIVE in (ws.parent / "out" / "d.md").read_text()
    summary = json.loads((ws.parent / "out" / "d.json").read_text())
    assert summary["review"]["occluded"] == [{"target": native.id, "page": 1, "quotes": [q.doc_text for q in item.quotes],
                                              "reason": "被悬浮的输入框盖住"}]


# ── Tables ──────────────────────────────────────────────────────────────


def test_the_agent_corrects_table_cells(ws):
    context = _context()
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    table = next(b for b in Workspace.open(ws).load().blocks if b.kind == BlockKind.TABLE)
    evidence = _evidence(ws, context, block=table.id)
    cell = {"op": "set_cells", "block": table.id, "reason": "图上是 8", "evidence": evidence,
            "cells": [{"row": 1, "col": 1, "content": "8"}]}
    assert _edit(ws, context, cell)[0].accepted
    grid = next(b for b in Workspace.open(ws).load().blocks if b.id == table.id).cells
    assert grid.slot(1, 1).content == "8" and grid.slot(1, 0).content == "SENTINEL-OCR 甲"
    outcome = _edit(ws, context, {**cell, "cells": [{"row": 9, "col": 0, "content": "x"}]})[0]
    assert not outcome.accepted and outcome.rule == "cell"


def test_an_empty_position_of_the_grid_may_be_filled(ws):
    from parserx.tables.grid import TableGrid

    context = _context()
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    workspace = Workspace.open(ws)
    with workspace.txn("test:missing cell") as state:  # (1, 1) was not recognized
        table = next(b for b in state.blocks if b.kind == BlockKind.TABLE)
        table.cells = TableGrid(n_rows=2, n_cols=2, cells=[c for c in table.cells.cells if (c.row, c.col) != (1, 1)])
    evidence = _evidence(ws, context, block=table.id)
    assert _edit(ws, context, {"op": "set_cells", "block": table.id, "reason": "图上是 12", "evidence": evidence,
                               "cells": [{"row": 1, "col": 1, "content": "12"}]})[0].accepted
    grid = next(b for b in Workspace.open(ws).load().blocks if b.id == table.id).cells
    assert grid.slot(1, 1).content == "12" and grid.slot(1, 0).content == "SENTINEL-OCR 甲"


def test_a_table_continued_across_pages_is_corrected_as_one(ws):
    # after the join the first block holds all rows and the anchors of every page; the continuation is merged
    from parserx.ir.anchor import PdfAnchor
    from parserx.ir.enums import BlockStatus
    from parserx.ir.relation import Relation

    context = _context()
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    workspace = Workspace.open(ws)
    with workspace.txn("test:merge") as state:
        second = next(b for b in state.blocks if b.kind == BlockKind.TABLE)
        first = second.model_copy(deep=True, update={
            "id": "t-first", "anchors": [PdfAnchor(page=1, bbox=(72, 500, 520, 700), coord_space="page_pt"),
                                         *second.anchors]})
        second.status = BlockStatus.MERGED
        state.blocks.append(first)
        state.relations.append(Relation(id="r-continues-t", kind="continues", src="t-first", dst=second.id))
    page2 = _evidence(ws, context, page=2, question="续表第 1 行的数值？", **{"as": "answer"})

    def cells(block, evidence, value):
        return _edit(ws, context, {"op": "set_cells", "block": block, "reason": f"图上是 {value}", "evidence": evidence,
                                   "cells": [{"row": 1, "col": 1, "content": value}]})[0]

    outcome = cells(second.id, page2, "8")
    assert not outcome.accepted and outcome.rule == "merged" and "t-first" in outcome.detail
    assert cells("t-first", page2, "8").accepted  # page 2 is one of the table's pages
    part = _evidence(ws, context, block=second.id, question="续表第 1 行的数值？", **{"as": "answer"})
    assert cells("t-first", part, "9").accepted  # the crop of a part merged into the table is evidence for the table


def test_look_at_the_seam_between_two_pages(ws):
    # a table or sentence continued on the next page: the bottom of one page and the top of the next in one image
    context = _context()
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    look = _look(ws, context, seam=1)
    assert look.seam == 1
    with Image.open(look.image.path) as image:
        assert image.height > image.width  # two half pages, stacked
    table = next(b for b in Workspace.open(ws).load().blocks if b.kind == BlockKind.TABLE)  # on page 2
    assert _edit(ws, context, {"op": "set_cells", "block": table.id, "reason": "图上是 8", "evidence": look.evidence,
                               "cells": [{"row": 1, "col": 1, "content": "8"}]})[0].accepted  # it shows both pages
    env, _ = _call("view_source", ws, {"looks": [{"seam": 2, "as": "answer", "question": "?"}]}, context=context)
    assert env.ok and env.result.results[0].evidence is None and env.failures[0].code == "invalid_request"  # no page 3


def test_look_at_some_rows_of_a_table_in_a_sharper_strip(ws):
    # a whole-table crop of a long table leaves its digits too small for the VLM (round 4, real_doc03_pdf):
    # looking at rows crops the band of those rows (with a row of margin) at a higher resolution
    context = _context()
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    table = next(b for b in Workspace.open(ws).load().blocks if b.kind == BlockKind.TABLE)
    whole = _look(ws, context, block=table.id)
    strip = _look(ws, context, block=table.id, rows=[1, 1])
    assert strip.image.asset != whole.image.asset and strip.image.width >= 1.9 * whole.image.width  # 300 dpi vs 150
    assert _edit(ws, context, {"op": "set_cells", "block": table.id, "reason": "图上是 8", "evidence": strip.evidence,
                               "cells": [{"row": 1, "col": 1, "content": "8"}]})[0].accepted
    env, _ = _call("view_source", ws, {"looks": [{"block": table.id, "rows": [5, 9]}]}, context=context)
    assert env.result.results[0].evidence is None and env.failures  # the table has 2 rows


# ── The VLM answers questions about the source ─────────────────────────


def test_an_answer_is_evidence_for_a_correction(ws):
    context = _context()
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    block = next(b for b in Workspace.open(ws).load().blocks if b.text == OCR_TEXT)
    env, code = _call("view_source", ws, {"looks": [{"block": block.id, "as": "answer",
                                                     "question": "这一行的数字是几？"}]}, context=context)
    data = _assert_contract(env, "view_source")
    answer = data["result"]["results"][0]
    assert code == 0 and data["cost"]["requests"] == {"vlm": 1} and answer["answer"]["doc_text"].startswith("SENTINEL-VLM")
    assert context.fake_vlm.calls[-1] == "parserx_ask_image"
    # the VLM's reading is evidence for a correction: the agent never looked itself
    assert _edit(ws, context, {"op": "replace_text", "block": block.id, "find": "3 件", "replace": "8 件",
                               "reason": "VLM 读图为 8 件", "evidence": answer["evidence"]})[0].accepted


def test_several_looks_in_one_call(ws):
    context = _context()
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    block = next(b for b in Workspace.open(ws).load().blocks if b.text == OCR_TEXT)
    looks = [{"block": block.id, "as": "answer", "question": "数字是几？"},
             {"page": 1, "as": "answer", "question": "有几张表？"},
             {"block": "b-missing", "as": "answer", "question": "？"},
             {"page": 1}]
    env, code = _call("view_source", ws, {"looks": looks}, context=context)
    data = _assert_contract(env, "view_source")
    assert env.ok and code == 0 and data["cost"]["requests"] == {"vlm": 2}
    results = data["result"]["results"]
    assert [r.get("block") for r in results] == [block.id, None, "b-missing", None] and results[1]["page"] == 1
    assert [bool(r.get("evidence")) for r in results] == [True, True, False, True] and results[3]["image"]
    assert [f["targets"] for f in data["failures"]] == [["b-missing"]] and data["failures"][0]["code"] == "not_found"


# ── Q42: text and tables inside embedded images become content after the image ──


def test_recognize_transcribes_an_embedded_image(ws, tmp_path):
    context = _context()
    figure = next(b for b in Workspace.open(ws).load().blocks if b.kind == BlockKind.FIGURE)
    env, code = _call("recognize", ws, {"blocks": [figure.id], "engine": "paddleocr"}, context=context)
    data = _assert_contract(env, "recognize")
    assert env.ok and code == 0 and data["cost"]["requests"] == {"ocr": 1}
    state = Workspace.open(ws).load()
    children = [r.dst for r in state.relations if r.kind == "contains" and r.src == figure.id]
    order = [b.id for b in sorted(state.blocks, key=lambda b: b.order)]
    at = order.index(figure.id)
    assert children and order[at + 1:at + 1 + len(children)] == children  # right after the image
    assert figure.id not in [b.id for b in state.blocks if b.status != "ok"]  # the image stays shown (Q42)
    assert verify_workspace(ws).ok
    accounts = _accounts(ws)
    assert accounts.accounting.unassigned == 0 and accounts.mismatched == []
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    env, _ = _call("export", ws, {"out": str(tmp_path / "out")}, context=context)
    markdown = (tmp_path / "out" / "doc.md").read_text()
    image_line = next(line for line in markdown.splitlines() if line.startswith("![") and figure.anchors[-1].asset in line)
    after = markdown[markdown.index(image_line):]
    assert after.index("<!-- parserx:image-text") < after.index("> **〔图片识别〕**") < after.index("SENTINEL-OCR 扫描文字") \
        < after.index("<!-- /parserx:image-text -->")  # not described: a picture, shown above its text (IO6-5)
    again, _ = _call("recognize", ws, {"blocks": [figure.id], "engine": "paddleocr"}, context=context)
    assert again.cost.requests == {} and again.failures and "already" in again.failures[0].message


def _routed(ws, route):
    from parserx.ir.state import ImageRecord

    workspace = Workspace.open(ws)
    figure = next(b for b in workspace.load().blocks if b.kind == BlockKind.FIGURE)
    with workspace.txn("test:route") as state:  # as the layout step would have routed it
        state.images = [ImageRecord(id=figure.anchors[-1].asset, route=route, shown=True, t=0.8, f=0.0, regions=3)]
    return figure


@pytest.mark.parametrize("kind, read", [("content", True), ("screenshot", False), ("photo", False)])
def test_the_pipeline_describes_first_and_transcribes_only_content(ws, kind, read):
    # IO6: the description says what the image is for; only an image whose words are its content is transcribed —
    # a picture is shown with its note, whatever its route
    from parserx.ir.enums import ImageRoute

    config = _config()
    config.runtime.layout_shadow = False
    context = _context()
    context.fake_vlm.figure_type = kind
    figure = _routed(ws, ImageRoute.SCAN)
    env, _ = _call("run_pipeline", ws, {}, config=config, context=context)
    assert env.ok and ("transcribe_images" in [s.step for s in env.result.steps]) is read
    state = Workspace.open(ws).load()
    assert next(b for b in state.blocks if b.id == figure.id).semantic.type == kind
    assert any(r.kind == "contains" and r.src == figure.id for r in state.relations) is read
    record = next(r for r in state.images if r.id == figure.anchors[-1].asset)
    assert (record.reading is not None) is read  # a content image's local reading checks its text (IO6-5)


@pytest.mark.parametrize("route, read", [("SCAN", True), ("MIXED", True), ("UNCERTAIN", False)])
def test_without_descriptions_images_are_transcribed_by_their_route(ws, route, read):
    # the service model off (or failing): the route decides, so no words are lost
    from parserx.ir.enums import ImageRoute

    config = _config()
    config.runtime.layout_shadow = False
    config.runtime.describe_figures = False
    figure = _routed(ws, ImageRoute(route))
    env, _ = _call("run_pipeline", ws, {}, config=config, context=_context())
    state = Workspace.open(ws).load()
    assert any(r.kind == "contains" and r.src == figure.id for r in state.relations) is read
    assert next(b for b in state.blocks if b.id == figure.id).semantic is None


def test_an_image_with_nothing_to_read_is_not_read_again(ws):
    # a photo: the engine finds only an image region, which folds back into the figure — no block follows it
    from parserx.ir.enums import ImageRoute
    from parserx.ir.state import ImageRecord

    def photo():
        return {"prunedResult": {"width": 1000, "height": 1400, "parsing_res_list": [
            {"block_label": "image", "block_content": "", "block_bbox": [0, 0, 1000, 1400], "block_order": None}]}}

    config = _config()
    config.runtime.layout_shadow = False
    context = _context(page=photo)
    context.fake_vlm.figure_type = "content"
    workspace = Workspace.open(ws)
    figure = next(b for b in workspace.load().blocks if b.kind == BlockKind.FIGURE)
    with workspace.txn("test:route") as state:
        state.images = [ImageRecord(id=figure.anchors[-1].asset, route=ImageRoute.MIXED, shown=True, t=0.3, f=0.2,
                                    regions=2)]
    first, _ = _call("run_pipeline", ws, {}, config=config, context=context)
    assert "transcribe_images" in [s.step for s in first.result.steps]
    again, _ = _call("run_pipeline", ws, {}, config=config, context=context)
    assert "transcribe_images" not in [s.step for s in again.result.steps] and again.cost.requests == {}
    assert again.result.check.exportable
    env, _ = _call("recognize", ws, {"blocks": [figure.id], "engine": "paddleocr"}, context=context)
    assert env.failures[0].message.endswith("already transcribed") and env.cost.requests == {}


def test_text_read_inside_an_image_can_be_looked_at_and_corrected(ws):
    # a block read inside an embedded image is anchored in the image's pixels: its crop comes from the image, and a
    # look at the whole image (the figure) is evidence for it, as a whole page is for a block on the page
    from parserx.ir.enums import ImageRoute
    from parserx.ir.state import ImageRecord

    config = _config()
    config.runtime.layout_shadow = False
    context = _context()
    context.fake_vlm.figure_type = "content"
    workspace = Workspace.open(ws)
    figure = next(b for b in workspace.load().blocks if b.kind == BlockKind.FIGURE)
    with workspace.txn("test:route") as state:
        state.images = [ImageRecord(id=figure.anchors[-1].asset, route=ImageRoute.SCAN, shown=True, t=0.8, f=0.0,
                                    regions=3)]
    _call("run_pipeline", ws, {}, config=config, context=context)
    state = Workspace.open(ws).load()
    inside = next(b for b in state.blocks if b.id.startswith(figure.id + "-") and b.text == OCR_TEXT)
    crop = _look(ws, context, block=inside.id)
    assert Path(crop.image.path).is_file() and crop.image.height < 100  # a line of the 128×100 image
    whole = _evidence(ws, context, block=figure.id, question="图中的件数是多少？", **{"as": "answer"})
    assert _edit(ws, context, {"op": "replace_text", "block": inside.id, "find": "3 件", "replace": "8 件",
                               "reason": "图上是 8 件", "evidence": whole})[0].accepted
    assert "8 件" in next(b for b in Workspace.open(ws).load().blocks if b.id == inside.id).text


def test_an_image_in_a_docx_is_transcribed_too(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("正文在图片之前")
    document.add_picture(io.BytesIO(_png(300, 200, (230, 230, 230))))
    document.add_paragraph("正文在图片之后")
    path = tmp_path / "scan.docx"
    document.save(path)
    workspace_init(path, tmp_path / "wsd", config=_config())
    context = _context()
    figure = next(b for b in Workspace.open(tmp_path / "wsd").load().blocks if b.kind == BlockKind.FIGURE)
    env, _ = _call("recognize", tmp_path / "wsd", {"blocks": [figure.id], "engine": "paddleocr"}, context=context)
    assert env.ok and env.cost.requests == {"ocr": 1}
    env, _ = _call("export", tmp_path / "wsd", {"out": str(tmp_path / "outd")}, context=context)
    markdown = (tmp_path / "outd" / "scan.md").read_text()
    assert markdown.index("正文在图片之前") < markdown.index("<!-- parserx:image-text src=") < \
        markdown.index("SENTINEL-OCR 扫描文字") < markdown.index("正文在图片之后")


def test_a_picture_in_a_table_cell_is_exported_with_the_table(ws):
    # the recognize step renders the page for a picture the engine left inside a cell (not only for figures)
    def page():
        cell = '<img src="imgs/img_in_image_box_100_300_300_400.jpg" alt="Image" /> 跨中'
        return {"prunedResult": {"width": 1000, "height": 1400, "parsing_res_list": [
            {"block_label": "table", "block_content": f"<table><tr><td>图示</td></tr><tr><td>{cell}</td></tr></table>",
             "block_bbox": [100, 280, 900, 500], "block_order": 1}]}}

    context = _context(page=page)
    _call("recognize", ws, {"pages": [2], "engine": "paddleocr"}, context=context)
    env, _ = _call("export", ws, {"out": str(ws.parent / "outp")}, context=context)
    md = Path(env.result.markdown).read_text()
    assert "img_in_image_box" not in md and "〔图1〕" in md
    assert md.index("〔图1〕") < md.index("<!-- 以下是上方〔图 n〕处的图片 -->") < md.rindex("](images/")


# ── Structure through edit_draft ────────────────────────────────────────


def test_structure_changes_return_the_issues_they_open(ws):
    # a structure left pending is open work: the call that made it says so (the agent sees what its change opened)
    block = next(b for b in json.loads((ws / "state.json").read_text())["blocks"] if b["kind"] == "text")["id"]
    env, _ = _call("edit_draft", ws, {"ops": [{"op": "mark_pending", "block": block, "reason": "test"}]})
    data = _assert_contract(env, "edit_draft")
    assert [(u["target"], u["kind"]) for u in data["result"]["issues_opened"]] == [(block, "structure_pending")]


def test_split_reports_the_new_block(tmp_path):
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 90), "3 Results\nThe measured values follow.", fontsize=11)
    doc.save(tmp_path / "two.pdf")
    ws = tmp_path / "ws"
    workspace_init(tmp_path / "two.pdf", ws, config=_config())
    block = next(b for b in json.loads((ws / "state.json").read_text())["blocks"] if "\n" in (b["text"] or ""))
    env, _ = _call("edit_draft", ws, {"ops": [{"op": "split", "block": block["id"], "at_break": 1,
                                               "reason": "title joined to the next line"}]})
    data = _assert_contract(env, "edit_draft")
    assert data["result"]["outcomes"][0]["accepted"] and data["result"]["outcomes"][0]["block"] == block["id"] + "-s1"
    assert any(b.id == block["id"] + "-s1" for b in Workspace.open(ws).load().blocks)


def test_without_a_scan_engine_the_document_still_exports_as_partial(ws):
    # `parserx parse --no-ocr`: scanned pages stay unrecognised, listed as missing; the run does not fail
    config = _config()
    config.builders.ocr.engine = "none"
    config.runtime.layout_shadow = False

    class NoScanEngine(_context()):
        def _new_ocr(self):
            return ToolContext._new_ocr(self)  # the real factory: refuses an engine that is not configured

    env, code = _call("run_pipeline", ws, {}, config=config, context=NoScanEngine)
    assert code == 0
    result = _assert_contract(env, "run_pipeline")["result"]
    assert result["pages"] == {"done": 1, "failed": 1} and result["check"]["exportable"]
    assert result["check"]["document_status"] == "partial"
    assert any("scan engine not configured" in f.message for f in env.failures)


def test_numbers_are_compared_with_the_reading_as_a_recognizer_confuses_them():
    from parserx.tools.describe_figure import unseen_numbers

    caption = "激光切割机，机身标有“LMN6000H”，额定 1,200 W，约 12 台，编号 7"
    assert unseen_numbers(caption, ["LMN6000H", "1200W"]) == []
    assert unseen_numbers(caption, ["LMN6O00H", "1 200 W"]) == []  # O read for 0, a space in a number
    assert unseen_numbers(caption, ["LMN60 0H"]) == ["6000", "1,200"]  # "约 12" is an estimate; 7 is one digit


@pytest.mark.parametrize("kind, seen, listed", [("photo", "型号 LMN6000H", False), ("photo", "型号 LMN60 0H", True),
                                                 ("content", "", False)])
def test_a_number_the_image_reading_lacks_is_open_work(ws, kind, seen, listed):
    # IO6-4: a picture's description quoting a number the local reading of the image does not have — the model may
    # have misread a digit — is listed until the description changes; a content image's text is transcribed instead
    from parserx.tools.describe_figure import Described, apply
    from parserx.tools.views import unresolved_items

    class Sees:
        name, version = "reading", "sees-1"

        def read(self, png):
            return [((0, 0, 10, 10), seen, 0.9)] if seen else []

    class Context(_context()):
        def _new_reader(self):
            return Sees()

    Context.fake_vlm.figure_type = kind
    Context.fake_vlm.caption = "激光切割机，机身标有“LMN6000H”，约 12 台"
    config = _config()
    config.runtime.layout_shadow = config.runtime.page_reading = False
    _call("run_pipeline", ws, {}, config=config, context=Context)
    state = Workspace.open(ws).load()
    items = [u for u in unresolved_items(state) if u.kind == "caption_number_unseen"]
    assert bool(items) is listed
    if listed:
        assert all([q.doc_text for q in u.quotes] == ["6000"] for u in items)
        for item in items:  # described again: the new description is what the item was about
            figure = next(b for b in state.blocks if b.id == item.target)
            apply(state, Described(figure.id, figure.anchors[-1], figure.semantic, None, ""), "vlm")
        assert not [u for u in unresolved_items(state) if u.kind == "caption_number_unseen"]

